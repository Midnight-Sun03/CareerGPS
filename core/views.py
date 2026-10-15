import datetime
import random
import json
import os
import uuid
from django.conf import settings
from django.http import HttpResponseNotAllowed, JsonResponse
from django.views.decorators.http import require_POST
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.contrib.auth import login, logout, authenticate
from django.contrib.auth.models import User
from django.contrib import messages
from django.core.paginator import Paginator
from .models import (
    Achievement, CareerPreference, Certification, Education, Interest, Language,
    Opportunity, Profile, ProfileFeedState, ProfileInterest, ProfileLink,
    ProfileOpportunityEvaluation, ProfileSkill, Project,
    SavedOpportunity, Skill, Story, MotivationalNudge, Volunteering,
    WorkExperience, Notification, LegalAcceptanceEvent,
)
from .legal import current_document
from .forms import (
    AchievementForm, CareerPreferenceForm, CertificationForm, EducationForm,
    LanguageForm, PersonalDetailsForm, ProfileInterestForm,
    ProfileLinkForm, ProfileSkillForm, ProjectForm, RegistrationForm, StoryForm,
    VolunteeringForm, WorkExperienceForm,
)
from .profile_projection import get_profile_projection, get_profile_progress
from .professional_matching import profile_has_signal
from .professional_policies import (
    FEED_CONFIRMED_CATEGORIES, FEED_DISCOVERY_CATEGORIES, TITLE_MATCH_CLASS,
    TITLE_MATCH_TIER,
)
from .services.title_gate import exclude_senior
from .matching import get_matched_opportunities
from .services.professional_evaluations import professional_feed_queryset
from .services.provisional_feed import provisional_evaluations
from .services.feed_v2 import (
    UNCERTAIN_CATALOGUE_LIMIT, authored_intent_fingerprint, feed_transition,
    has_authored_targets, proved_unrelated_opportunity_ids,
    unassessed_catalogue_queryset, v2_feed_queryset,
)
from .services.career_intent import profile_careers
from .services.feed_explanations import explain_sparse_feed, transition_disclosure
from .services.notifications import (
    CLOSING_SOON_DAYS, notify_overdue_opportunities_for_user, published_closing_date,
)
from applications.models import Application, GeneratedDocument
from applications.services.candidate_facts import build_candidate_fact_bundle
from applications.services.document_generation import (
    assemble_legacy_cv, generate_cover_letter as generate_controlled_cover_letter,
    generate_cv_writing,
)
from applications.services.job_context import job_context_text, resolve_job_context
from applications.services.json_resume import build_json_resume
from applications.services.semantic_relevance import select_relevant_facts
from applications.services.document_timing import document_timing, timed_stage
from applications.services.resume_templates import (
    normalize_layout_settings, normalize_template_key, public_template_options, render_resume_pdf,
)
import secrets
import time


def calculate_completion(profile):
    return get_profile_progress(profile)[0]


def get_progress_items(profile):
    return get_profile_progress(profile)[1]

from django.core.mail import EmailMultiAlternatives
from django.db import transaction
from django.db.models import Q
from django.utils.crypto import constant_time_compare, salted_hmac
import logging

logger = logging.getLogger(__name__)

def _extract_text_from_cv(uploaded_file):
    """Takes a Django UploadedFile and returns its plain text content."""
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        import pdfplumber

        text_parts = []
        with pdfplumber.open(uploaded_file) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n".join(text_parts)

    elif filename.endswith(".docx"):
        from docx import Document

        doc = Document(uploaded_file)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    else:
        raise ValueError("Unsupported file type. Please upload a .pdf or .docx file.")
 
def landing(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    return render(request, 'core/landing.html')


OTP_EXPIRY_SECONDS = 10 * 60
OTP_RESEND_COOLDOWN_SECONDS = 60
OTP_MAX_ATTEMPTS = 5
OTP_MAX_RESENDS = 5


def _otp_digest(user_id, otp):
    return salted_hmac(
        "careergps.email-verification",
        f"{user_id}:{otp}",
    ).hexdigest()


def _clear_otp_session(request, *, keep_user=False):
    keys = [
        "otp_hash",
        "otp_expires_at",
        "otp_attempts",
        "otp_last_sent_at",
        "otp_resend_count",
    ]

    for key in keys:
        request.session.pop(
            key,
            None,
        )

    if not keep_user:
        request.session.pop(
            "user_id",
            None,
        )

    request.session.modified = True


def _store_otp_session(
    request,
    user,
    otp,
    *,
    resend_count=0,
):
    now = int(
        time.time()
    )

    request.session[
        "user_id"
    ] = user.id

    request.session[
        "otp_hash"
    ] = _otp_digest(
        user.id,
        otp,
    )

    request.session[
        "otp_expires_at"
    ] = (
        now
        + OTP_EXPIRY_SECONDS
    )

    request.session[
        "otp_attempts"
    ] = 0

    request.session[
        "otp_last_sent_at"
    ] = now

    request.session[
        "otp_resend_count"
    ] = resend_count

    request.session.modified = True


def _send_verification_email(
    user,
    otp,
):
    """
    Send a CareerGPS verification OTP to the user's
    registered email address.

    If SMTP is not configured or delivery fails, an
    exception is raised and handled by the registration/
    resend view. The OTP is never printed to the terminal.
    """

    if not getattr(
        settings,
        "EMAIL_HOST",
        "",
    ):
        raise RuntimeError(
            "CareerGPS email delivery is not configured."
        )

    subject = (
        "Your CareerGPS verification code"
    )

    text_body = (
        "Welcome to CareerGPS.\n\n"
        f"Your verification code is: {otp}\n\n"
        "This code expires in 10 minutes.\n\n"
        "If you did not create a CareerGPS account, "
        "you can ignore this email."
    )

    html_body = f"""
    <div style="
        font-family: Arial, sans-serif;
        max-width: 560px;
        margin: 0 auto;
        color: #222;
    ">
        <h2 style="
            margin-bottom: 8px;
        ">
            Verify your CareerGPS email
        </h2>

        <p>
            Use the verification code below to
            complete your CareerGPS registration.
        </p>

        <div style="
            font-size: 30px;
            font-weight: 700;
            letter-spacing: 8px;
            padding: 18px 20px;
            margin: 24px 0;
            background: #f7f7f7;
            border-radius: 10px;
            text-align: center;
        ">
            {otp}
        </div>

        <p>
            This code expires in
            <strong>10 minutes</strong>.
        </p>

        <p style="
            color: #666;
            font-size: 13px;
            margin-top: 30px;
        ">
            If you did not create a CareerGPS
            account, you can ignore this email.
        </p>
    </div>
    """

    email = EmailMultiAlternatives(
        subject=subject,
        body=text_body,
        from_email=(
            settings.DEFAULT_FROM_EMAIL
        ),
        to=[
            user.email
        ],
    )

    email.attach_alternative(
        html_body,
        "text/html",
    )

    sent = email.send(
        fail_silently=False
    )

    if sent != 1:
        raise RuntimeError(
            "CareerGPS could not send the verification email."
        )


def _create_and_send_otp(
    request,
    user,
    *,
    resend_count=0,
):
    otp = str(
        secrets.randbelow(
            900000
        )
        + 100000
    )

    # Send first.
    #
    # We only replace the session OTP after the email
    # has actually been accepted by the configured
    # email backend.
    _send_verification_email(
        user,
        otp,
    )

    _store_otp_session(
        request,
        user,
        otp,
        resend_count=resend_count,
    )


def register_view(request):
    if request.user.is_authenticated:
        return redirect(
            "dashboard"
        )

    if request.method == "POST":
        form = RegistrationForm(
            request.POST
        )

        if form.is_valid():
            try:
                with transaction.atomic():
                    email = (
                        form.cleaned_data["email"]
                        .strip()
                        .lower()
                    )

                    # Recover an existing unverified registration
                    # instead of attempting to create a duplicate
                    # Django username/account.
                    user = (
                        User.objects
                        .filter(
                            email__iexact=email,
                            is_active=False,
                        )
                        .order_by("id")
                        .first()
                    )

                    if user is None:
                        user = form.save()

                        Profile.objects.create(
                            user=user
                        )

                    else:
                        user.first_name = (
                            form.cleaned_data[
                                "first_name"
                            ].strip()
                        )

                        user.last_name = (
                            form.cleaned_data[
                                "last_name"
                            ].strip()
                        )

                        user.email = email
                        user.username = email

                        user.set_password(
                            form.cleaned_data[
                                "password"
                            ]
                        )

                        user.save(
                            update_fields=[
                                "first_name",
                                "last_name",
                                "email",
                                "username",
                                "password",
                            ]
                        )

                        Profile.objects.get_or_create(
                            user=user
                        )

                    for kind, action in (
                        ('terms', LegalAcceptanceEvent.Action.AGREED),
                        ('privacy', LegalAcceptanceEvent.Action.ACKNOWLEDGED),
                    ):
                        document = current_document(kind)
                        LegalAcceptanceEvent.objects.get_or_create(
                            user=user, document_type=kind,
                            document_version=document.version,
                            content_sha256=document.content_sha256,
                            action=action, context='registration',
                        )

                    _create_and_send_otp(
                        request,
                        user,
                    )

            except Exception:
                logger.exception(
                    "CareerGPS registration verification delivery failed."
                )
                _clear_otp_session(
                    request
                )

                form.add_error(
                    None,
                    (
                        "CareerGPS could not send your "
                        "verification email right now. "
                        "Please check your email address "
                        "and try again."
                    ),
                )

            else:
                return redirect(
                    "verify_otp"
                )

    else:
        form = RegistrationForm()

    return render(
        request,
        "core/register.html",
        {
            "form": form
        },
    )

def verify_otp(request):
    user_id = request.session.get(
        "user_id"
    )

    if not user_id:
        messages.error(
            request,
            (
                "Your verification session has expired. "
                "Please register again."
            ),
        )

        return redirect(
            "register"
        )

    try:
        user = User.objects.get(
            id=user_id
        )

    except User.DoesNotExist:
        _clear_otp_session(
            request
        )

        messages.error(
            request,
            (
                "Your verification session is no "
                "longer valid. Please register again."
            ),
        )

        return redirect(
            "register"
        )

    if user.is_active:
        _clear_otp_session(
            request
        )

        return redirect(
            "login"
        )

    if request.method == "POST":
        entered_otp = str(
            request.POST.get(
                "otp",
                ""
            )
        ).strip()

        stored_hash = request.session.get(
            "otp_hash"
        )

        expires_at = int(
            request.session.get(
                "otp_expires_at",
                0,
            )
            or 0
        )

        attempts = int(
            request.session.get(
                "otp_attempts",
                0,
            )
            or 0
        )

        # --------------------------------------------
        # No active OTP
        # --------------------------------------------

        if not stored_hash:
            messages.error(
                request,
                (
                    "There is no active verification "
                    "code. Please request a new code."
                ),
            )

            return render(
                request,
                "core/verify_otp.html",
                {
                    "email": user.email,
                },
            )

        # --------------------------------------------
        # Expired
        # --------------------------------------------

        if (
            not expires_at
            or int(time.time())
            > expires_at
        ):
            _clear_otp_session(
                request,
                keep_user=True,
            )

            messages.error(
                request,
                (
                    "That verification code has expired. "
                    "Please request a new code."
                ),
            )

            return render(
                request,
                "core/verify_otp.html",
                {
                    "email": user.email,
                },
            )

        # --------------------------------------------
        # Attempt limit
        # --------------------------------------------

        if attempts >= OTP_MAX_ATTEMPTS:
            _clear_otp_session(
                request,
                keep_user=True,
            )

            messages.error(
                request,
                (
                    "Too many incorrect attempts. "
                    "Please request a new verification code."
                ),
            )

            return render(
                request,
                "core/verify_otp.html",
                {
                    "email": user.email,
                },
            )

        # --------------------------------------------
        # Validate format
        # --------------------------------------------

        if (
            len(entered_otp) != 6
            or not entered_otp.isdigit()
        ):
            messages.error(
                request,
                (
                    "Enter the 6-digit code "
                    "from your email."
                ),
            )

            return render(
                request,
                "core/verify_otp.html",
                {
                    "email": user.email,
                },
            )

        entered_hash = _otp_digest(
            user.id,
            entered_otp,
        )

        # --------------------------------------------
        # Correct OTP
        # --------------------------------------------

        if constant_time_compare(
            stored_hash,
            entered_hash,
        ):
            user.is_active = True

            user.save(
                update_fields=[
                    "is_active"
                ]
            )

            _clear_otp_session(
                request
            )

            login(
                request,
                user,
            )

            messages.success(
                request,
                (
                    f"Welcome to CareerGPS, "
                    f"{user.first_name}! "
                    "Complete your profile to see "
                    "personalised opportunities."
                ),
            )

            return redirect(
                "dashboard"
            )

        # --------------------------------------------
        # Wrong OTP
        # --------------------------------------------

        attempts += 1

        request.session[
            "otp_attempts"
        ] = attempts

        request.session.modified = True

        attempts_left = (
            OTP_MAX_ATTEMPTS
            - attempts
        )

        if attempts_left <= 0:
            _clear_otp_session(
                request,
                keep_user=True,
            )

            messages.error(
                request,
                (
                    "Too many incorrect attempts. "
                    "Please request a new verification code."
                ),
            )

        else:
            messages.error(
                request,
                (
                    "Incorrect verification code. "
                    f"{attempts_left} attempt"
                    f"{'' if attempts_left == 1 else 's'} "
                    "remaining."
                ),
            )

    return render(
        request,
        "core/verify_otp.html",
        {
            "email": user.email,
        },
    )


@require_POST
def resend_otp(request):
    user_id = request.session.get(
        "user_id"
    )

    if not user_id:
        messages.error(
            request,
            (
                "Your verification session has expired. "
                "Please register again."
            ),
        )

        return redirect(
            "register"
        )

    try:
        user = User.objects.get(
            id=user_id,
            is_active=False,
        )

    except User.DoesNotExist:
        _clear_otp_session(
            request
        )

        messages.error(
            request,
            (
                "Your verification session is no "
                "longer valid. Please register again."
            ),
        )

        return redirect(
            "register"
        )

    now = int(
        time.time()
    )

    last_sent = int(
        request.session.get(
            "otp_last_sent_at",
            0,
        )
        or 0
    )

    resend_count = int(
        request.session.get(
            "otp_resend_count",
            0,
        )
        or 0
    )

    # --------------------------------------------
    # 60-second resend cooldown
    # --------------------------------------------

    elapsed = (
        now
        - last_sent
    )

    if (
        last_sent
        and elapsed
        < OTP_RESEND_COOLDOWN_SECONDS
    ):
        wait_seconds = (
            OTP_RESEND_COOLDOWN_SECONDS
            - elapsed
        )

        messages.error(
            request,
            (
                "Please wait "
                f"{wait_seconds} seconds before "
                "requesting another code."
            ),
        )

        return redirect(
            "verify_otp"
        )

    # --------------------------------------------
    # Resend limit
    # --------------------------------------------

    if resend_count >= OTP_MAX_RESENDS:
        messages.error(
            request,
            (
                "You have requested too many "
                "verification codes. Please register "
                "again later."
            ),
        )

        return redirect(
            "verify_otp"
        )

    try:
        _create_and_send_otp(
            request,
            user,
            resend_count=(
                resend_count + 1
            ),
        )

    except Exception:
        messages.error(
            request,
            (
                "CareerGPS could not send another "
                "verification email right now. "
                "Please try again shortly."
            ),
        )

        return redirect(
            "verify_otp"
        )

    messages.success(
        request,
        (
            "A new verification code has been "
            f"sent to {user.email}."
        ),
    )

    return redirect(
        "verify_otp"
    )


def login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    
    error = None
    
    if request.method == 'POST':
        email = request.POST.get('email')
        password = request.POST.get('password')
        
        try:
            user = User.objects.get(email=email)
            user = authenticate(
                request,
                username=user.username,
                password=password
            )
            if user:
                login(request, user)
                return redirect('dashboard')
            else:
                error = "Invalid email or password."
        except User.DoesNotExist:
            error = "No account found with that email."
    
    return render(request, 'core/login.html', {'error': error})

@login_required
def logout_confirm(request):
    return render(
        request,
        "core/logout_confirm.html",
    )


@login_required
@require_POST
def logout_view(request):
    logout(request)

    messages.success(
        request,
        "You have been logged out successfully.",
    )

    return redirect(
        "landing"
    )

def _tier_zero_cards(profile):
    """Cards for the dashboard's tier-zero band: undecided plus proven-unrelated.

    ``unassessed_catalogue_queryset`` intentionally excludes jobs whose stored
    evidence proves them unrelated for every authored target.  The dashboard
    still needs those rows for profile completeness, especially when the only
    catalogue job is a proven non-match, so collect both sets independently
    and combine them.  Neither set is part of the confirmed feed.
    """
    scope = _candidate_universe()
    proved_unrelated = proved_unrelated_opportunity_ids(
        profile, list(
            scope.values_list('active_clean_revision_id', flat=True)))
    undecided = list(unassessed_catalogue_queryset(
        profile, scope).select_related('active_clean_revision'))
    undecided_ids = {
        opportunity.active_clean_revision_id for opportunity in undecided}
    unrelated = list(Opportunity.objects.filter(
        active_clean_revision_id__in=set(proved_unrelated) - undecided_ids,
    ).select_related('active_clean_revision').order_by(
        '-active_clean_revision__posted_sort_at', 'pk'))
    return (
        [_uncertain_card(opportunity) for opportunity in undecided]
        + [_uncertain_card(opportunity, proved_unrelated=True)
           for opportunity in unrelated])


@login_required
def dashboard_view(request):
    profile = get_object_or_404(Profile, user=request.user)
    notify_overdue_opportunities_for_user(request.user)
    saved_count = SavedOpportunity.objects.filter(
        user=request.user
    ).count()
    today = timezone.localdate()
    week_from_now = today + datetime.timedelta(days=CLOSING_SOON_DAYS)
    closing_soon = SavedOpportunity.objects.filter(
        user=request.user,
        opportunity__in=Opportunity.objects.professionally_ready(),
        opportunity__active_clean_revision__closing_sort_at__date__range=(today, week_from_now),
    ).count()
    completion_percentage = calculate_completion(profile)
    authoritative_ready = _authoritative_feed_ready(profile)
    # Prepared current-identity legacy results are only consulted when no
    # authoritative generation matches the current intent, so an ACTIVE v2 feed
    # always wins.  None means "not ready yet" and falls through to the truthful
    # catalogue fallback below.
    provisional_rows = None if authoritative_ready else (
        _prepared_provisional_rows(profile))
    if authoritative_ready:
        # Authoritative personalised feed is current for this target.  Cards
        # are backed by stored v2 admissions; the runtime matcher only supplies
        # current-fit evidence text for those admitted opportunities, so this
        # request never re-runs matching over the whole catalogue.
        admitted_ids = list(v2_feed_queryset(profile).values_list(
            'opportunity_id', flat=True)[:25])
        if not admitted_ids:
            # ACTIVE generation matches the current intent, but the confirmed
            # feed is empty.  Surface the undecided + proven-unrelated
            # catalogue as the dashboard's tier-zero band, so jobs that are
            # still undecided or have been proved unrelated stay visible on
            # the dashboard while remaining excluded from the confirmed feed.
            matched_results = _tier_zero_cards(profile)[:DASHBOARD_PROVISIONAL_LIMIT]
        else:
            ranked = get_matched_opportunities(
                profile, queryset=_candidate_universe().filter(pk__in=admitted_ids),
                include_cv_tips=False) if admitted_ids else []
            ranked_index = {item['opportunity'].pk: item for item in ranked}
            matched_results = [
                _admission_card(row, ranked_index, confirmed=True)
                for row in v2_feed_queryset(profile)[:5]]
    elif provisional_rows:
        # Target-specific matched and ranked results already prepared by the
        # background provisional task. No matcher runs on this request.
        matched_results = [
            _provisional_card(row) for row in provisional_rows[:5]]
    else:
        # No ACTIVE generation matches the current intent: either the profile has
        # no authored targets yet, or the target changed and the new generation
        # has not activated.  Both are honestly a catalogue, not a match list, so
        # the dashboard reads the provisional population directly instead of
        # running the expensive matcher.  Ordering is the existing provisional
        # ordering (most recently posted first), which carries no claim of
        # evidence-ranked relevance.
        matched_results = [
            _uncertain_card(opportunity)
            for opportunity in unassessed_catalogue_queryset(
                profile, _candidate_universe()).select_related(
                    'active_clean_revision')[:DASHBOARD_PROVISIONAL_LIMIT]
        ]
    saved_ids = set(SavedOpportunity.objects.filter(
        user=request.user
    ).values_list("opportunity_id", flat=True))
    for result in matched_results:
        result["is_saved"] = result["opportunity"].id in saved_ids
    applications = {row.opportunity_id: row for row in Application.objects.filter(
        user=request.user,
        opportunity_id__in=[result["opportunity"].id for result in matched_results],
    ).only("id", "opportunity_id", "status")}
    for result in matched_results:
        application = applications.get(result["opportunity"].id)
        result["has_application"] = application is not None
        result["application_action_label"] = application.workspace_action_label if application else "Prepare"
    top_matches = matched_results[:5]
    feed_explanation = None
    if not top_matches:
        feed_explanation = {
            'title': 'No current recommendations',
            'message': (
                'No eligible opportunities are available right now. New '
                'listings are validated automatically as they arrive.'),
        }

    nudge_trigger = "profile_incomplete" if completion_percentage < 60 else "general"
    nudge = MotivationalNudge.objects.filter(
        trigger=nudge_trigger, is_active=True
    ).order_by("?").first()
    if not nudge:
        nudge = MotivationalNudge.objects.filter(trigger="general", is_active=True).order_by("?").first()

    v2_enabled = settings.CAREERGPS_FEED_READ_VERSION == 'v2'
    # Career-intelligence presentation boundary: the dashboard reuses the same
    # stored evidence the Career Intelligence page reads, with no new queries.
    # Disclosure fields stay raw on the context and are converted to candidate
    # language only in the template, so the matching/admission verdict is never
    # altered and nothing internal reaches the rendered output.
    # Career targets are authored as free text on CareerPreference.target_roles
    # (comma / semicolon / newline separated); the model has no target_ordinal
    # column, which is what the previous code wrongly assumed.  ``has_authored_
    # targets`` is the feed's single completeness authority for "does this
    # profile have targets", and the authored target list mirrors the same rule
    # Career Intelligence uses (``profile_careers`` -> ``parse_target_roles``),
    # enumerating ordinals over the authored text itself.  Reading targets this
    # way keeps the dashboard in step with a changed career target immediately,
    # without depending on the derived resolution rows and without touching the
    # matching engine.
    has_targets = has_authored_targets(profile)
    authored_targets = [
        {"label": text, "ordinal": ordinal}
        for ordinal, text in enumerate(profile_careers(profile))
    ]
    selected_target = authored_targets[0] if authored_targets else None
    selected_ordinal = selected_target["ordinal"] if selected_target else None
    employer_counts = {}
    for title, disc, count in _matched_employers(top_matches, profile):
        employer_counts[title] = (disc, count)
    analysis = {
        "has_targets": has_targets,
        "selected": selected_target,
        "selected_ordinal": selected_ordinal,
        "authored_targets": authored_targets,
        "matched": {
            "employers": [
                [title, employer_counts.get(title, ("", 0)), count]
                for title, count in sorted(
                    employer_counts.items(), key=lambda kv: -kv[1][1])
            ]
        },
        "evidence": {
            "buckets": _evidence_buckets(profile),
        },
    }
    context = {
        "profile": profile,
        "saved_count": saved_count,
        "closing_soon": closing_soon,
        "completion_percentage": completion_percentage,
        "top_matches": top_matches,
        "feed_explanation": feed_explanation,
        "matches_today": len(top_matches),
        "has_professional_signal": (
            (authoritative_ready and bool(v2_feed_queryset(profile).exists()))
            or bool(provisional_rows)
            if v2_enabled
            else profile_has_signal(profile)),
        "show_profile_prompt": completion_percentage < 60,
        "nudge": nudge,
        "has_targets": has_targets,
        "analysis": analysis,
    }
    return render(request, "core/dashboard.html", context)


def _clean_location(revision):
    value = revision.location or {}
    if isinstance(value, str):
        return value
    parts = [str(value.get(key) or "").strip() for key in (
        "city", "region", "country") if value.get(key)]
    if parts:
        return ", ".join(parts)
    # Parsed city/region/country may be empty while the original scraped
    # location text is populated; show that rather than discarding it.
    raw = str(value.get("raw") or "").strip()
    return raw or "Location not stated"


def _evidence_buckets(profile):
    """Plain-language evidence buckets reused by the dashboard header.

    Presentation only: buckets are derived from profile progress items that
    already exist on the request.  The raw evidence text is never displayed
    here -- the template emits the short candidate-facing label instead.
    """
    buckets = [
        {"label": "Profile ready", "count": 0, "items": []},
        {"label": "Not sure yet", "count": 0, "items": []},
        {"label": "Not there yet", "count": 0, "items": []},
    ]
    try:
        for item in get_progress_items(profile):
            if item.get("state") == "shown":
                buckets[0]["count"] += 1
            elif item.get("state") == "cannot_confirm":
                buckets[1]["count"] += 1
            else:
                buckets[2]["count"] += 1
    except Exception:
        pass
    return buckets


def _matched_employers(top_matches, profile):
    """Yield (title, disclosure, count) candidates from the top match cards.

    The disclosure is intentionally LEFT raw here: the presentation boundary
    is enforced in the template, where raw internal wording is replaced with
    candidate-facing copy.
    """
    seen = set()
    for m in top_matches:
        rev = m.get("clean_revision")
        title = (getattr(rev, "company", None)
                 or getattr(rev, "cleaned_title", None)
                 or "") if rev is not None else ""
        if not title or title in seen:
            continue
        seen.add(title)
        yield (title, "", 1)


PROVISIONAL_PAGE_SIZE = 24
# The dashboard has always shown exactly five cards. The provisional
# experience keeps that same slice, so the layout and the "5" statistic are
# unchanged; only what those five cards honestly claim changes.
DASHBOARD_PROVISIONAL_LIMIT = 5


def _authoritative_feed_ready(profile):
    """True only when an ACTIVE v2 generation matches the current intent.

    This is the single switch that separates the authoritative personalised feed
    from the provisional catalogue.  It deliberately mirrors the two conditions
    `v2_feed_queryset` itself enforces: the generation must be ACTIVE, and its
    stored intent hash must still equal the intent derived from the current
    authored targets.  A target change therefore returns False immediately,
    before any new generation exists, which is what makes the old feed
    disappear instead of being shown against the new target.
    """
    state = ProfileFeedState.objects.select_related('active_generation').filter(
        profile=profile).first()
    if not state or not state.active_generation_id:
        return False
    generation = state.active_generation
    return bool(
        generation.state == 'ACTIVE'
        and generation.intent_dependency_hash == authored_intent_fingerprint(profile)
    )


# ------------------------------------------------------- provisional bridge
#
# Precedence for the non-authoritative state is:
#   current-identity prepared legacy results > truthful catalogue fallback.
# `_authoritative_feed_ready` is always checked first by the callers, so an
# ACTIVE v2 feed can never be displaced by these rows.

_BAND_PRIORITY = {'A': 6, 'B': 5, 'C': 4, 'D': 3, 'E': 2, 'F': 1}
_PROVISIONAL_TIER = {'TARGET': 3, 'RELATED': 2, 'PROFILE_RELEVANT': 0}
_PROVISIONAL_LABEL = {3: 'Exact target-role alignment',
                      2: 'Related target-role alignment',
                      0: 'No supported target-role alignment'}


def _prepared_provisional_rows(profile):
    """Prepared current-identity legacy results, or None when none are usable.

    Returns None rather than an empty list so callers can fall through to the
    existing truthful catalogue fallback without a second condition.
    """
    if not has_authored_targets(profile):
        return None
    rows = provisional_evaluations(
        profile, ProfileOpportunityEvaluation.objects.filter(
            profile=profile).select_related('opportunity', 'clean_revision'))
    if not rows.exists():
        return None
    return sorted(rows, key=_ranking_key_of)


def _ranking_key_of(row):
    """The legacy mandatory order, read from the stored row."""
    return (
        -_PROVISIONAL_TIER.get(row.discovery_category, 0),
        -_BAND_PRIORITY.get(row.band, 1),
        -int(row.professional_score),
        row.pk,
    )


def _provisional_card(evaluation):
    """Present one prepared legacy row on the existing candidate card contract.

    Field names match what `_candidate_card` already reads, so no template or
    card-contract change is needed and the old ranking labels are preserved.
    """
    tier = _PROVISIONAL_TIER.get(evaluation.discovery_category, 0)
    reasons = evaluation.discovery_reasons or []
    return _candidate_card({
        'opportunity': evaluation.opportunity,
        'clean_revision': evaluation.clean_revision,
        'score': int(evaluation.professional_score),
        'recommendation_band': evaluation.band,
        'recommendation_priority': _BAND_PRIORITY.get(evaluation.band, 1),
        'recommendation_label': _PROVISIONAL_LABEL.get(
            tier, 'No supported target-role alignment'),
        'career_direction_tier_priority': tier,
        'career_direction_label': _PROVISIONAL_LABEL.get(
            tier, 'No supported target-role alignment'),
        'career_direction_tier_label': _PROVISIONAL_LABEL.get(
            tier, 'No supported target-role alignment'),
        'why_text': str(reasons[0]) if reasons else '',
        'major_gap_count': evaluation.major_gap_count,
        'gate_gap_count': evaluation.gate_gap_count,
    })


def _candidate_universe():
    """Shared candidate feed universe for every candidate.

    Normal lifecycle/readiness rules stay authoritative, and the one shared
    read-only senior decision is applied on top: the same `evaluate_revision`
    decision, including the cleaner's description-based evidence, that Career
    Intelligence uses.  Nothing is read from or written to a persisted
    seniority row, so the two pages cannot disagree about a senior role and a
    policy change needs no backfill.
    """
    return exclude_senior(Opportunity.objects.professionally_ready(
        include_obvious_senior=False))


_RELATIONSHIP_STATE = {
    'EXACT': ('Exact occupational match',
              'Confirmed by ESCO or OFO as the same occupation as {scope}.'),
    'RELATED': ('Approved related role',
                'Confirmed by ESCO or OFO as an approved related role for {scope}.'),
    'CONFLICT': ('Conflicting evidence',
                 'Occupational evidence for {scope} is conflicting.'),
    'UNRESOLVED': (
        'Relationship not established',
        'Occupation relationship not established for {scope}.'),
    TITLE_MATCH_CLASS: (
        'Title match only',
        'The title matches {scope} but ESCO and OFO have not confirmed the '
        'occupation relationship.'),
}


# A discovery-area card is never labelled a confirmed match. The title-only
# class gets its own wording so title resemblance is never presented as
# taxonomy confirmation.
_DISCOVERY_RELATIONSHIP_LABEL = {
    TITLE_MATCH_CLASS: (
        'Title match only',
        'The title resembles {scope}, but neither ESCO nor OFO confirmed the '
        'occupation relationship, so this is discovery only.'),
}


# How one taxonomy phrased the link between two occupations, when it gave one
# at all.  Silence has no wording here, because silence is not a relationship.
_RELATIONSHIP_VERBS = {
    'EXACT': 'is the same occupation as',
    'RELATED': 'is related to',
}

# Retained title-only evidence carries its own tier label, so a possibility is
# never rendered in the wording of a confirmation.
_TITLE_ONLY_TIER_LABELS = {
    TITLE_MATCH_TIER: 'Title-only evidence, not taxonomy-confirmed',
}

# The uncertain area has no current personalised ranking, because these jobs
# have no current assessment, so the page says how they are ordered instead of
# implying a score.  Confirmed and assessed jobs keep the stored global ranking
# in their own areas.
UNCERTAIN_ORDERING_NOTE = (
    'Ordered by how recently the listing was posted. These jobs have no stored '
    'assessment against your targets yet, so no fit score or match ranking is '
    'shown for them and none is reused from an earlier feed.')


def _relationship_side_name(label, code):
    """One occupation of one side, named and coded so two stays two.

    ESCO and OFO can both match the bare word "accountant" for different
    concepts.  The concept code is what tells those two occupations apart, so it
    belongs in the displayed name and in the identity used to recognise a
    duplicate, not in a field nobody renders.
    """
    name = ' '.join(str(label or '').split())
    code = ' '.join(str(code or '').split())
    if name and code:
        return f'{name} ({code})'
    return name or code


def _relationship_lines(evaluation):
    """Every occupational link a stored admission established, as readable lines.

    The card carries one category because a card is one job. These are the
    individual ESCO and OFO relationships behind that category: a job that ESCO
    links to three target occupations and OFO to a fourth is shown as four
    relationships, and nothing is merged into a single sentence.

    Two lines are the same relationship only when they share their taxonomy and
    both concept codes.  Relationships are therefore never merged on the
    strength of matching display text, which two different occupations can
    easily produce.
    """
    if not hasattr(evaluation, 'admission_reasons'):
        return []
    lines = []
    seen = set()
    for row in evaluation.admission_reasons or []:
        for link in row.get('relationships') or []:
            verb = _RELATIONSHIP_VERBS.get(link.get('class'))
            if verb is None:
                continue
            target_code = str(link.get('target_code') or '').strip()
            job_code = str(link.get('job_code') or '').strip()
            # The concept ids carry the identity of the relationship.  Codes and
            # labels are display text, and two distinct occupations in one
            # taxonomy can share a label and have no code at all, so they must
            # never be merged into one line on the strength of that text.
            identity = (
                (link.get('source') or ''),
                (row.get('assessment_id'), row.get('target_resolution_id')),
                link.get('target_concept_id'),
                link.get('job_concept_id'),
                target_code,
                str(link.get('target_label') or '').strip(),
                job_code,
                str(link.get('job_label') or '').strip(),
                str(link.get('relationship_type') or ''),
            )
            if identity in seen:
                continue
            seen.add(identity)
            lines.append({
                'source': (link.get('source') or '').upper(),
                'authored_target': row.get('authored_target') or 'your target',
                'target_occupation': _relationship_side_name(
                    link.get('target_label'), target_code),
                'target_code': target_code,
                'verb': verb,
                'job_occupation': _relationship_side_name(
                    link.get('job_label'), job_code),
                'job_code': job_code,
                'relationship_type': (link.get('relationship_type') or '')
                                     .replace('_', ' ').lower(),
            })
    return lines


def _professional_card(opportunity, evaluation=None, *, fit_updating=False):
    revision = opportunity.active_clean_revision
    result = {
        "opportunity": opportunity,
        "clean_revision": revision,
        "location": _clean_location(revision),
        "discovery_label": "Opportunity catalogue",
        "fit_label": "Not yet evaluated",
        "card_reason": "No professional-fit conclusion is available for this opportunity.",
        "major_warning": "",
        "relationship_reason": "",
        "relationship_label": "",
        "relationship_lines": [],
        "title_only_lines": [],
    }
    if evaluation:
        result.update({
            "evaluation": evaluation,
            "discovery_label": evaluation.get_discovery_category_display(),
            "fit_label": (
                "Current fit updating" if fit_updating
                else f"Band {evaluation.band}"),
            "card_reason": (
                f"{evaluation.positive_important_count} important criteria have positive "
                f"profile evidence; {evaluation.major_gap_count} major gaps are recorded."
            ),
            "major_warning": (
                f"{evaluation.gate_gap_count} gate criteria are not found on the profile."
                if evaluation.gate_gap_count else ""
            ),
        })
        if hasattr(evaluation, 'admission_reasons'):
            targets = [
                row.get('authored_target') for row in evaluation.admission_reasons
                if row.get('authored_target')]
            scope = (', '.join(dict.fromkeys(targets)) if targets
                     else 'your career target')
            label, message = _RELATIONSHIP_STATE.get(
                evaluation.discovery_category,
                ('Relationship not established',
                 'Occupation relationship not established for {scope}.'))
            result['relationship_label'] = label
            result['relationship_reason'] = message.format(scope=scope)
            result['relationship_lines'] = _relationship_lines(evaluation)
    return result


_TIER_DISPLAY_LABELS = {
    3: "Exact match",
    2: "Related role",
    1: "Adjacent role",
    0: "General opportunity",
}


def _candidate_card(item):
    """Present one runtime match item on the shared candidate cards.

    Presentation adapter only: band, tier, evidence and score fields are
    copied straight from the deterministic matcher. No classification is
    invented here and nothing here can remove an opportunity.
    """
    opportunity = item["opportunity"]
    revision = opportunity.active_clean_revision
    tier = item.get("career_direction_tier_priority", 0)
    card = dict(item)
    card.update({
        "opportunity": opportunity,
        "clean_revision": revision,
        "location": (
            _clean_location(revision) if revision else "Location not stated"),
        "discovery_label": _TIER_DISPLAY_LABELS.get(
            tier, "General opportunity"),
        "fit_label": (
            f'{item.get("recommendation_label", "")} '
            f'(Band {item.get("recommendation_band", "?")})'),
        "card_reason": item.get("why_text", ""),
        "major_warning": (
            f'{item["major_gap_count"]} major required gap(s) are recorded.'
            if item.get("major_gap_count") else ""),
        "relationship_label": (
            item.get("career_direction_label", "") if tier else ""),
        "relationship_reason": (
            item.get("career_direction_tier_label", "") if tier else ""),
        "is_saved": False,
        "has_application": False,
        "application_action_label": "Prepare",
    })
    return card


def _kind_filter(value):
    # Entry-level is not a revision kind: it resolves against the structured
    # `Opportunity.opportunity_type` classification inside the feed branches.
    return {
        "internship": "INTERNSHIP", "learnership": "LEARNERSHIP",
        "graduate": "GRADUATE_PROGRAMME",
    }.get(value)


def _title_only_lines(evaluation):
    """Retained title-only evidence of a stored admission, as readable lines.

    A policy-free same-concept label collision was really compared and is kept,
    but it established nothing.  It is disclosed on its own lines, each with the
    tier that marks it as not taxonomy-confirmed, so title resemblance can never
    be read as one of the confirmed relationships listed above it.
    """
    if not hasattr(evaluation, 'admission_reasons'):
        return []
    lines = []
    seen = set()
    for row in evaluation.admission_reasons or []:
        for item in row.get('title_only_evidence') or []:
            identity = (
                (item.get('source') or ''),
                (row.get('assessment_id'), row.get('target_resolution_id')),
                # Concept ids identify the evidence; two distinct occupations
                # that merely share a label are two pieces of evidence.
                item.get('target_concept_id'),
                item.get('job_concept_id'),
                item.get('target_code'),
                item.get('job_code'),
                item.get('evidence_tier'),
            )
            if identity in seen:
                continue
            seen.add(identity)
            tier = item.get('evidence_tier') or TITLE_MATCH_TIER
            lines.append({
                'source': (item.get('source') or '').upper(),
                'target_occupation': _relationship_side_name(
                    item.get('target_label'), item.get('target_code')),
                'job_occupation': _relationship_side_name(
                    item.get('job_label'), item.get('job_code')),
                'tier': tier,
                'tier_label': _TITLE_ONLY_TIER_LABELS.get(
                    tier, 'Title-only evidence'),
                'reason_code': item.get('reason_code') or '',
            })
    return lines


def _uncertain_card(opportunity, ranked=None, *, proved_unrelated=False):
    """One catalogue job that is not part of the confirmed feed.

    Presentation of a stored fact only: an undecided listing has not been
    assessed against the candidate's targets, while a proven-unrelated one
    carries fresh stored evidence that it matches none of them.  Neither is a
    confirmed match, no stored personalised score is reused here, and the
    legacy ranking tier is deliberately ignored, because a job with no current
    occupational admission has no current personalised verdict to quote.
    """
    revision = opportunity.active_clean_revision
    card = dict(ranked or {})
    card.update({
        "opportunity": opportunity,
        "clean_revision": revision,
        "location": _clean_location(revision),
        "discovery_label": (
            "Proven unrelated" if proved_unrelated else "Not yet assessed"),
        "relationship_label": (
            "Proven unrelated" if proved_unrelated else "Not yet assessed"),
        "relationship_reason": (
            "Stored occupational evidence proves this listing does not match "
            "any of your current targets, so it is kept visible here but is "
            "excluded from your confirmed recommendations."
            if proved_unrelated else
            "CareerGPS has not yet assessed this listing against your targets, "
            "so it is not presented as a confirmed match. It is also not proved "
            "unrelated: until the assessment exists, this job is simply "
            "undecided."),
        "relationship_lines": [],
        "title_only_lines": [],
        "is_confirmed_relationship": False,
        "admission_class": "",
        "card_reason": (
            "Current fit is measured from this listing's requirements and your "
            "profile. Occupational relevance is not assessed yet."
            if not proved_unrelated else
            "Shown for profile completeness. Occupational evidence rules out a "
            "current target match."),
        "major_warning": "",
        "fit_label": "Current fit updating",
        "is_saved": False,
        "has_application": False,
        "application_action_label": "Prepare",
    })
    return card


def _admission_card(evaluation, ranked_index, *, confirmed):
    """Present one stored v2 admission using the existing card contract.

    This is presentation only.  The occupational verdict is copied from the
    stored assessment; nothing here can add, remove or reclassify an
    opportunity.  A card in the discovery area is never labelled an exact or
    related match, even when the ranking-token tier would have said so.
    """
    opportunity = evaluation.opportunity
    revision = evaluation.clean_revision
    ranked = ranked_index.get(opportunity.pk, {})
    card = dict(ranked)
    card.update({
        "opportunity": opportunity,
        "clean_revision": revision,
        "location": _clean_location(revision),
        "is_confirmed_relationship": confirmed,
        "admission_class": evaluation.discovery_category,
        "relationship_tier": evaluation.relationship_tier,
        "relationship_lines": _relationship_lines(evaluation),
        "title_only_lines": _title_only_lines(evaluation),
        "fit_label": (
            f'{ranked.get("recommendation_label", "")} '
            f'(Band {evaluation.band})'),
        "card_reason": (
            f"{evaluation.positive_important_count} important criteria have "
            f"positive profile evidence; {evaluation.major_gap_count} major "
            "gaps are recorded."),
        "major_warning": (
            f"{evaluation.gate_gap_count} gate criteria are not found on the "
            "profile." if evaluation.gate_gap_count else ""),
        "is_saved": False,
        "has_application": False,
        "application_action_label": "Prepare",
    })
    if confirmed:
        card["discovery_label"] = (
            "Exact match" if evaluation.discovery_category == 'EXACT'
            else "Related role")
        targets = [
            row.get('authored_target') for row in evaluation.admission_reasons
            if row.get('authored_target')]
        scope = ', '.join(dict.fromkeys(targets)) or 'your career target'
        label, message = _RELATIONSHIP_STATE.get(
            evaluation.discovery_category,
            ('Relationship not established',
             'Occupation relationship not established for {scope}.'))
        card["relationship_label"] = label
        card["relationship_reason"] = message.format(scope=scope)
    else:
        # Never "Exact match" / "Related role" from a ranking token alone.
        label, message = _DISCOVERY_RELATIONSHIP_LABEL.get(
            evaluation.discovery_category,
            ('Not yet confirmed',
             'CareerGPS has not yet confirmed an occupation match for this '
             'listing against your targets. It is not proved unrelated, so it '
             'is listed here for discovery only.'))
        card["discovery_label"] = label
        card["relationship_label"] = label
        card["relationship_reason"] = message.format(
            scope=scope_for_card(evaluation))
    return card


def scope_for_card(evaluation):
    """The authored targets that produced a stored admission, for display."""
    targets = [
        row.get('authored_target') for row in evaluation.admission_reasons
        if row.get('authored_target')]
    if targets:
        return ', '.join(dict.fromkeys(targets))
    return 'your career target'


@login_required
def opportunities_view(request):
    profile = get_object_or_404(Profile, user=request.user)
    notify_overdue_opportunities_for_user(request.user)
    search = request.GET.get("search", "").strip()
    opportunity_type = request.GET.get("type", "all")
    if opportunity_type not in {"all", "internship", "learnership",
                                "graduate", "entry_level"}:
        opportunity_type = "all"
    kind = _kind_filter(opportunity_type)
    entry_level_only = opportunity_type == "entry_level"

    # Layers 1-2: the shared catalogue and the shared read-only senior check.
    scope = _candidate_universe()
    if kind:
        scope = scope.filter(active_clean_revision__opportunity_kind=kind)
    if entry_level_only:
        scope = scope.filter(opportunity_type="entry_level")
    if search:
        scope = scope.filter(
            Q(active_clean_revision__cleaned_title__icontains=search)
            | Q(active_clean_revision__company__icontains=search)
            | Q(active_clean_revision__cleaned_description__icontains=search)
        )

    has_targets = has_authored_targets(profile)
    # One switch decides which experience the page is offering.  It is True
    # whenever no ACTIVE generation matches the current intent, so a target
    # change flips it before the replacement generation exists.
    provisional_primary = has_targets and not _authoritative_feed_ready(profile)
    # Prepared current-identity legacy results, consulted only while v2 is not
    # authoritative for the current intent.
    prepared_provisional = None
    if provisional_primary:
        prepared_provisional = _prepared_provisional_rows(profile)
    uses_provisional_feed = bool(prepared_provisional)

    # Layer 3: occupational relevance, read only from stored v2 admissions.
    # No relationship is ever derived on this request.
    confirmed_ids = set()
    discovery_ids = set()
    uncertain_ids = set()
    uncertain = []
    uncertain_count = 0
    uncertain_page_obj = None
    provisional_page_obj = None
    if has_targets:
        # Search, type and entry-level filters run inside the same universe for
        # every area: a practical filter changes what is on screen, never what
        # the stored occupational admission concluded.
        confirmed_ids = set(v2_feed_queryset(
            profile, categories=FEED_CONFIRMED_CATEGORIES
        ).filter(opportunity__in=scope).values_list('opportunity_id', flat=True))
        discovery_ids = set(v2_feed_queryset(
            profile, categories=FEED_DISCOVERY_CATEGORIES
        ).filter(opportunity__in=scope).values_list('opportunity_id', flat=True))
        # Requirements and eligibility never remove a job here: the stored
        # admission already decided relevance, and the matcher is used only
        # for card evidence and the "why" text on jobs the stored feed
        # already admitted.  It never matches the whole catalogue at runtime,
        # and it does not run at all while a provisional feed is serving the
        # page (provisional rows carry their own prepared cards).
        admitted_ids = set(confirmed_ids) | set(discovery_ids)
        ranked = []
        if not uses_provisional_feed and admitted_ids:
            ranked = get_matched_opportunities(
                profile, queryset=scope.filter(pk__in=list(admitted_ids)),
                include_cv_tips=False)
        ranked_index = {item["opportunity"].pk: item for item in ranked}
        evaluations = list(v2_feed_queryset(
            profile, categories=FEED_CONFIRMED_CATEGORIES).filter(
                opportunity__in=scope))
        paginator = Paginator(evaluations, 24)
        page_obj = paginator.get_page(request.GET.get("page"))
        matched_results = [
            _admission_card(row, ranked_index, confirmed=True)
            for row in page_obj.object_list]
        # The discovery area never repeats a confirmed job, so one job is
        # always in exactly one assessed area.
        discovery = [
            _admission_card(row, ranked_index, confirmed=False)
            for row in v2_feed_queryset(
                profile, categories=FEED_DISCOVERY_CATEGORIES
            ).filter(opportunity__in=scope).exclude(
                opportunity_id__in=confirmed_ids)[:60]
        ]
        # Unassessed catalogue jobs stay browsable in their own area, next to
        # the real transition message, instead of the page looking empty while
        # occupational coverage is missing or stale.  Stored data only: the
        # senior-filtered `scope` is reused as it is, and nothing is resolved,
        # compared or written here.  A job that the active generation already
        # admitted, or that fresh evidence proves UNRELATED for every authored
        # target, is not in this set at all.
        uncertain_scope = unassessed_catalogue_queryset(profile, scope)
        uncertain_count = uncertain_scope.count()
        uncertain_scope = uncertain_scope.select_related('active_clean_revision')
        # While no authoritative generation is current this catalogue is the
        # primary experience, so it is browsable in full on the same page size
        # the confirmed area already uses rather than a fixed peek.  Once a
        # generation is authoritative the uncertain area is genuinely
        # supplemental and keeps its capped, truncated-disclosure behaviour.
        if provisional_primary:
            uncertain_paginator = Paginator(uncertain_scope, PROVISIONAL_PAGE_SIZE)
            uncertain_page_obj = uncertain_paginator.get_page(
                request.GET.get("page"))
            # The confirmed area is empty in this state, so the page's single
            # existing pagination control drives the provisional catalogue
            # instead of an empty confirmed list.
            page_obj = uncertain_page_obj
            uncertain = [
                _uncertain_card(opportunity)
                for opportunity in uncertain_page_obj.object_list]
        else:
            uncertain_page_obj = None
            uncertain = [
                _uncertain_card(opportunity, ranked_index.get(opportunity.pk))
                for opportunity in uncertain_scope[:UNCERTAIN_CATALOGUE_LIMIT]]
        uncertain_ids = {
            row["opportunity"].pk for row in uncertain}
    else:
        # No usable target roles: an honestly labelled general feed, unchanged
        # in shape, with a prompt to add targets.  There is no target, so there
        # is nothing to assess and nothing for the matcher to rank against a
        # direction; the existing provisional catalogue ordering is used.
        provisional_scope = unassessed_catalogue_queryset(
            profile, scope).select_related('active_clean_revision')
        provisional_paginator = Paginator(provisional_scope, PROVISIONAL_PAGE_SIZE)
        provisional_page_obj = provisional_paginator.get_page(
            request.GET.get("page"))
        # The confirmed-area pagination control is the page's only one, so the
        # provisional catalogue is paged through it rather than adding a second
        # navigation element to the page.
        page_obj = provisional_page_obj
        matched_results = [
            _uncertain_card(opportunity)
            for opportunity in provisional_page_obj.object_list]
        discovery = []

    if uses_provisional_feed:
        # Target-specific matched + ranked rows already prepared in the
        # background. They replace the empty confirmed area so the candidate
        # sees useful matches immediately, and they carry no matcher call on
        # this request. Paginated with the page's existing control.
        provisional_paginator = Paginator(prepared_provisional,
                                          PROVISIONAL_PAGE_SIZE)
        provisional_page_obj = provisional_paginator.get_page(
            request.GET.get("page"))
        matched_results = [_provisional_card(row)
                           for row in provisional_page_obj.object_list]
        page_obj = provisional_page_obj
        # The catalogue fallback is redundant once prepared matches exist.
        uncertain = []
        uncertain_count = 0
        uncertain_page_obj = None

    cards = matched_results + discovery + uncertain
    saved_ids = set(
    SavedOpportunity.objects.filter(
        user=request.user
    ).values_list("opportunity_id", flat=True))

    for result in cards:
         result["is_saved"] = result["opportunity"].id in saved_ids
    applications = {row.opportunity_id: row for row in Application.objects.filter(
        user=request.user,
        opportunity_id__in=[result["opportunity"].id for result in cards],
    ).only("id", "opportunity_id", "status")}
    for result in cards:
        application = applications.get(result["opportunity"].id)
        result["has_application"] = application is not None
        result["application_action_label"] = application.workspace_action_label if application else "Prepare"

    feed_explanation = None
    if not matched_results:
        if search:
            feed_explanation = {
                'title': 'No recommendations match this search',
                'message': (
                    'Search runs inside the current opportunity universe, '
                    'so no separate catalogue is opened.'),
                'action_url': reverse('opportunities'),
                'action_label': 'Clear search',
            }
        elif has_targets:
            feed_explanation = {
                'title': 'No confirmed matches for your targets yet',
                'message': (
                    'No current opportunity is confirmed as an exact or '
                    'approved related match to your saved targets. This is not '
                    'an empty market: CareerGPS is still resolving occupations, '
                    'so unconfirmed listings appear in the separate section '
                    'below and jobs with no assessment yet appear in their own '
                    'section under it.'),
            }
        else:
            feed_explanation = {
                'title': 'Showing all current opportunities',
                'message': (
                    'Add target roles in Profile -> Career Preferences to see '
                    'which opportunities are confirmed matches for your goals.'),
                'action_url': reverse('profile_section',
                                      kwargs={'section': 'career-preferences'}),
                'action_label': 'Add your targets',
            }

    # Disclosure-only: the stored transition state already decided what the
    # feed may show; this surfaces why, so an updating feed is never read as a
    # smaller market.  Nothing here adds, removes or re-ranks a card.
    feed_transition_note = transition_disclosure(feed_transition(profile))

    context = {
        "results": matched_results,
        "discovery_results": discovery,
        # A clearly separate uncertain area: eligible catalogue jobs that no
        # fresh stored assessment covers.  Its ordering is stated on the page,
        # because these jobs carry no current personalised ranking.
        "uncertain_results": uncertain,
        "uncertain_count": uncertain_count,
        "uncertain_ids": uncertain_ids,
        "uncertain_ordering": UNCERTAIN_ORDERING_NOTE,
        "page_obj": page_obj,
        "search": search,
        "selected_type": opportunity_type,
        "selected_feed": "all",
        "feed_explanation": feed_explanation,
        "feed_transition": feed_transition_note,
        "personalized_feed": has_targets,
        "has_targets": has_targets,
        "confirmed_count": len(confirmed_ids),
        "discovery_count": len(discovery_ids),
        # True while no ACTIVE generation matches the current intent: the
        # uncertain catalogue is then the primary experience and is browsable
        # in full, not a capped supplement.
        "provisional_primary": provisional_primary,
        "uses_provisional_feed": uses_provisional_feed,
        "uncertain_page_obj": uncertain_page_obj,
        "provisional_page_obj": provisional_page_obj,
    }

    return render(request, "core/opportunities.html", context)

@login_required
def opportunity_detail_view(request, opportunity_id):

    # Direct catalogue detail remains available for professionally ready rows;
    # personalized-feed seniority admission is a separate concern.
    detail_catalogue = Opportunity.objects.professionally_ready(
        include_obvious_senior=True)
    opportunity = get_object_or_404(
        detail_catalogue.select_related(
            "active_clean_revision"
        ),
        id=opportunity_id,
    )

    profile = get_object_or_404(
        Profile,
        user=request.user
    )

    is_saved = SavedOpportunity.objects.filter(
        user=request.user,
        opportunity=opportunity
    ).exists()
    application = Application.objects.filter(
        user=request.user, opportunity=opportunity
    ).first()
    has_application = application is not None

    evaluation = professional_feed_queryset(profile).filter(
        opportunity=opportunity,
        clean_revision=opportunity.active_clean_revision,
    ).first()
    match_data = _professional_card(opportunity, evaluation) if evaluation else None

    context = {
        "opportunity": opportunity,
        "clean_revision": opportunity.active_clean_revision,
        "location": (_clean_location(opportunity.active_clean_revision)
                     if opportunity.active_clean_revision else "Location not stated"),
        "profile": profile,
        "match_data": match_data,
        "is_saved": is_saved,
        "has_application": has_application,
        "application_action_label": application.workspace_action_label if application else "Prepare",
    }

    return render( request, "core/opportunity_detail.html", context)

PROFILE_SECTION_CONFIG = {
    'education': (Education, EducationForm, 'Education', 'This helps us find opportunities that match your education.'),
    'experience': (WorkExperience, WorkExperienceForm, 'Work Experience', 'Add paid, part-time, internship, or other relevant work.'),
    'projects': (Project, ProjectForm, 'Projects', 'Projects remain separate evidence of what you have built or contributed to.'),
    'skills': (ProfileSkill, ProfileSkillForm, 'Skills', 'Type any skill that genuinely belongs on your profile.'),
    'certifications': (Certification, CertificationForm, 'Certifications', 'Add credentials and expiry information where applicable.'),
    'achievements': (Achievement, AchievementForm, 'Achievements', 'Record awards or achievements that matter to your career story.'),
    'volunteering': (Volunteering, VolunteeringForm, 'Volunteering', 'Include community or volunteer experience relevant to your goals.'),
    'languages': (Language, LanguageForm, 'Languages', 'Add languages and your current proficiency.'),
    'interests': (ProfileInterest, ProfileInterestForm, 'Interests', 'Type career and industry interests in your own words.'),
    'links': (ProfileLink, ProfileLinkForm, 'Links & Portfolio', 'Share relevant professional work, profiles, or portfolios.'),
    'career-preferences': (
        CareerPreference, CareerPreferenceForm, 'Career Preferences',
        'Target roles drive personalized recommendations; practical preferences remain user-controlled filters.'),
}


def _display_date(value):
    return value.strftime('%b %Y') if value else ''


def _display_date_range(item, current_field):
    start = _display_date(getattr(item, 'start_date', None))
    is_current = bool(getattr(item, current_field, False))
    end = 'Present' if is_current else _display_date(getattr(item, 'end_date', None))
    return ' – '.join(value for value in (start, end) if value)


def _profile_entry_display(section, item):
    """Return concise, user-facing text instead of exposing model __str__ output."""
    title = ''
    meta_parts = []
    detail = ''

    if section == 'education':
        title = item.qualification or item.field_of_study or 'Education'
        meta_parts = [item.institution, item.field_of_study]
        detail = _display_date_range(item, 'currently_studying')
    elif section == 'experience':
        title = item.title
        meta_parts = [item.organisation, item.location]
        detail = _display_date_range(item, 'currently_working')
    elif section == 'projects':
        title = item.title
        meta_parts = [item.organisation]
        detail = _display_date_range(item, 'currently_ongoing')
    elif section == 'skills':
        title = item.entered_name
    elif section == 'certifications':
        title = item.title
        meta_parts = [item.issuer]
        start = _display_date(item.issue_date)
        end = 'No expiry' if item.does_not_expire else _display_date(item.expiry_date)
        detail = ' – '.join(value for value in (start, end) if value)
    elif section == 'achievements':
        title = item.title
        meta_parts = [item.organisation]
        detail = _display_date(item.date)
    elif section == 'volunteering':
        title = item.role
        meta_parts = [item.organisation, item.location]
        detail = _display_date_range(item, 'currently_volunteering')
    elif section == 'languages':
        title = item.language
        meta_parts = [item.get_proficiency_display()]
    elif section == 'links':
        title = item.label
        meta_parts = [item.get_link_type_display()]
        detail = item.url
    elif section == 'interests':
        title = item.entered_name
    elif section == 'career-preferences':
        title = item.target_roles.strip() or 'Career preferences'
        meta_parts = [item.preferred_locations.strip(), item.get_work_arrangement_display()]
        detail = item.short_career_goal.strip()

    return {
        'title': title or 'Untitled entry',
        'meta': ' • '.join(value for value in meta_parts if value),
        'detail': detail,
    }


def _section_hub_summary(slug, items):
    count = len(items)
    if not items:
        return 'Not added yet'
    if slug in {'skills', 'interests'}:
        names = [item.entered_name for item in items[:3]]
        return ', '.join(names) + ('…' if count > 3 else '')
    if slug == 'languages':
        names = [item.language for item in items[:3]]
        return ', '.join(names) + ('…' if count > 3 else '')
    return _profile_entry_display(slug, items[0])['title']


def _profile_hub_context(profile):
    tiles = []
    personal_complete = bool(
        profile.user.first_name or profile.user.last_name or profile.phone or profile.location
    )
    personal_summary = profile.user.get_full_name().strip() or profile.user.email
    if profile.location:
        personal_summary = f'{personal_summary} • {profile.location}'
    tiles.append({
        'slug': 'personal', 'title': 'Personal Info', 'count': int(personal_complete),
        'summary': personal_summary, 'complete': personal_complete, 'personal': True,
    })

    for slug, (model, _, title, _) in PROFILE_SECTION_CONFIG.items():
        items = list(model.objects.filter(profile=profile))
        tiles.append({
            'slug': slug,
            'title': title,
            'count': len(items),
            'summary': _section_hub_summary(slug, items),
            'complete': bool(items),
            'personal': False,
        })

    return {
        'profile': profile,
        'tiles': tiles,
        'completion_percentage': calculate_completion(profile),
        'progress_items': get_progress_items(profile),
    }


def _profile_section_context(profile, section, bound_form=None, edit_pk=None):
    model, form_class, title, helper_text = _profile_section_or_404(section)
    entries = []
    for item in model.objects.filter(profile=profile):
        form = form_class(instance=item, profile=profile, prefix=f'{section}-{item.pk}')
        if edit_pk == item.pk and bound_form is not None:
            form = bound_form
        entries.append({
            'object': item,
            'form': form,
            'display': _profile_entry_display(section, item),
        })

    add_form = form_class(profile=profile, prefix=f'add-{section}')
    if edit_pk is None and bound_form is not None:
        add_form = bound_form
    return {
        'profile': profile,
        'section': section,
        'section_title': title,
        'helper_text': helper_text,
        'entries': entries,
        'add_form': add_form,
        'single': model is CareerPreference,
        'skill_suggestions': Skill.objects.order_by('name') if section == 'skills' else (),
        'interest_suggestions': Interest.objects.order_by('name') if section == 'interests' else (),
        'completion_percentage': calculate_completion(profile),
    }


def _personal_section_context(profile, form=None):
    return {
        'profile': profile,
        'personal': True,
        'section': 'personal',
        'section_title': 'Personal Info',
        'helper_text': 'Keep your name and career contact details up to date.',
        'personal_form': form or PersonalDetailsForm(user=profile.user, profile=profile),
        'completion_percentage': calculate_completion(profile),
    }


@login_required
def profile_view(request):
    profile = get_object_or_404(Profile, user=request.user)
    return render(request, 'core/profile.html', _profile_hub_context(profile))


@login_required
def profile_complete_view(request):
    """Compatibility route for old bookmarks; legacy-only POSTs are disabled."""
    if request.method != 'GET':
        return HttpResponseNotAllowed(['GET'])
    return redirect('profile')


@login_required
def profile_personal_view(request):
    profile = get_object_or_404(Profile, user=request.user)
    return render(request, 'core/profile_section.html', _personal_section_context(profile))


@login_required
def profile_section_view(request, section):
    profile = get_object_or_404(Profile, user=request.user)
    return render(
        request,
        'core/profile_section.html',
        _profile_section_context(profile, section),
    )


@login_required
@require_POST
def update_personal_details(request):
    profile = get_object_or_404(Profile, user=request.user)
    form = PersonalDetailsForm(request.POST, user=request.user, profile=profile)
    if form.is_valid():
        form.save()
        messages.success(request, 'Personal details saved.')
        return redirect('profile_personal')
    return render(
        request, 'core/profile_section.html',
        _personal_section_context(profile, form=form), status=400,
    )


def _profile_section_or_404(section):
    config = PROFILE_SECTION_CONFIG.get(section)
    if config is None:
        from django.http import Http404
        raise Http404('Unknown profile section.')
    return config


@login_required
@require_POST
def add_profile_section(request, section):
    profile = get_object_or_404(Profile, user=request.user)
    model, form_class, title, _ = _profile_section_or_404(section)
    if model is CareerPreference and model.objects.filter(profile=profile).exists():
        messages.info(request, 'Edit your existing career preferences instead.')
        return redirect('profile_section', section=section)
    form = form_class(request.POST, profile=profile, prefix=f'add-{section}')
    if form.is_valid():
        form.save()
        messages.success(request, f'{title} saved.')
        return redirect('profile_section', section=section)
    return render(
        request, 'core/profile_section.html',
        _profile_section_context(profile, section, bound_form=form), status=400,
    )


@login_required
@require_POST
def edit_profile_section(request, section, record_id):
    profile = get_object_or_404(Profile, user=request.user)
    model, form_class, title, _ = _profile_section_or_404(section)
    instance = get_object_or_404(model, pk=record_id, profile=profile)
    form = form_class(
        request.POST, instance=instance, profile=profile, prefix=f'{section}-{record_id}'
    )
    if form.is_valid():
        form.save()
        messages.success(request, f'{title} updated.')
        return redirect('profile_section', section=section)
    return render(
        request, 'core/profile_section.html',
        _profile_section_context(
            profile, section, bound_form=form, edit_pk=record_id
        ), status=400,
    )


@login_required
@require_POST
def delete_profile_section(request, section, record_id):
    profile = get_object_or_404(Profile, user=request.user)
    model, _, title, _ = _profile_section_or_404(section)
    instance = get_object_or_404(model, pk=record_id, profile=profile)
    instance.delete()
    messages.info(request, f'{title} entry deleted.')
    return redirect('profile_section', section=section)

@login_required
def toggle_save_view(request, opportunity_id):
    """
    Saves or removes an opportunity from the user's saved list.
    """

    existing = SavedOpportunity.objects.select_related('opportunity').filter(
        user=request.user, opportunity_id=opportunity_id
    ).first()
    if existing:
        revision = existing.opportunity.active_clean_revision
        title = revision.cleaned_title if revision else "this opportunity"
        existing.delete()
        messages.info(
            request,
            f'"{title}" has been removed from your saved opportunities.'
        )
    else:
        opportunity = get_object_or_404(Opportunity.objects.professionally_ready(), id=opportunity_id)
        SavedOpportunity.objects.create(user=request.user, opportunity=opportunity)
        messages.success(
            request,
            f'"{opportunity.active_clean_revision.cleaned_title}" has been added to your saved opportunities.'
        )

    return redirect( request.META.get("HTTP_REFERER", "dashboard"))

@login_required
def saved_view(request):
    """
    Displays all opportunities saved by the current user.
    """

    from django.db.models import Prefetch
    from ingestion.models import OpportunitySourceRecord
    saved_opportunities = (
        SavedOpportunity.objects
        .filter(user=request.user)
        .select_related("opportunity", "opportunity__active_clean_revision")
        .prefetch_related(
            "opportunity__required_skills", "opportunity__preferred_skills",
            Prefetch(
                "opportunity__source_records",
                queryset=OpportunitySourceRecord.objects.select_related(
                    "source", "source__greenhouse_board"
                ).order_by("source__priority", "-last_seen_at"),
                to_attr="_application_source_records",
            )
        )
        .order_by("-saved_at")
    )
    profile = get_object_or_404(Profile, user=request.user)
    saved_opportunities = list(saved_opportunities)
    applications = {row.opportunity_id: row for row in Application.objects.filter(
        user=request.user,
        opportunity_id__in=[saved.opportunity_id for saved in saved_opportunities],
    ).select_related('clean_revision').only(
        'id', 'opportunity_id', 'status', 'clean_revision_id',
        'clean_revision__cleaned_title', 'clean_revision__company',
        'clean_revision__location', 'clean_revision__closing_at',
        'clean_revision__source_url', 'clean_revision__opportunity_kind')}
    evaluations = {row.opportunity_id: row for row in professional_feed_queryset(
        profile).filter(opportunity_id__in=[saved.opportunity_id for saved in saved_opportunities])}
    strong_matches = 0
    today = timezone.localdate()
    week_from_now = today + datetime.timedelta(days=CLOSING_SOON_DAYS)
    closing_soon = 0
    for saved in saved_opportunities:
        application = applications.get(saved.opportunity_id)
        saved.catalogue_available = Opportunity.objects.professionally_ready().filter(
            pk=saved.opportunity_id).exists()
        # Saved history may display its last immutable clean facts after the
        # listing closes. It never re-enters matching/catalogue admission.
        saved.clean_revision = (
            application.clean_revision if application and application.clean_revision_id
            else saved.opportunity.active_clean_revision)
        # Present a single readable location string (never the raw stored dict).
        saved.display_location = (
            _clean_location(saved.clean_revision) if saved.clean_revision
            else "Location not stated")
        evaluation = evaluations.get(saved.opportunity_id)
        saved.match_data = _professional_card(saved.opportunity, evaluation) if evaluation else None
        saved.has_application = application is not None
        saved.application_id = application.pk if application else None
        saved.application_action_label = application.workspace_action_label if application else "Prepare application"
        saved.opportunity_closed = (
            not saved.opportunity.is_active or saved.opportunity.lifecycle_status != 'active'
            or any(date and date < today for date in (
                saved.opportunity.closing_date, saved.opportunity.application_deadline
            ))
        )
        strong_matches += int(bool(evaluation and evaluation.band in {'A', 'B'}))
        deadline = (saved.opportunity.active_clean_revision.closing_sort_at.date()
                    if saved.catalogue_available and saved.opportunity.active_clean_revision.closing_sort_at
                    else None)
        closing_soon += int(bool(
            saved.catalogue_available and deadline and today <= deadline <= week_from_now
        ))

    return render(request, "core/saved.html", {
            "saved_opportunities": saved_opportunities,
            "saved_count": len(saved_opportunities),
            "strong_matches": strong_matches,
            "closing_soon": closing_soon,
        })

@login_required
def stories_view(request):
    """
    Display approved community stories.
    """

    stories = (
        Story.objects
        .filter(status="approved")
        .select_related("user")
        .order_by("-submitted_at")
    )

    context = {
        "stories": stories,
    }

    return render(request, "core/stories.html", context)

@login_required
def notifications_view(request):
    notify_overdue_opportunities_for_user(request.user)
    notifications = (
        Notification.objects
        .filter(user=request.user)
        .order_by(
            "-created_at",
            "-id",
        )
    )

    selected_notification = None

    selected_id = request.GET.get(
        "selected"
    )

    if selected_id:
        try:
            selected_notification = (
                Notification.objects.get(
                    id=int(selected_id),
                    user=request.user,
                )
            )

        except (
            ValueError,
            TypeError,
            Notification.DoesNotExist,
        ):
            selected_notification = None

    unread_count = (
        notifications
        .filter(is_read=False)
        .count()
    )
    page = Paginator(notifications, 30).get_page(request.GET.get('page'))
    page.object_list = list(page.object_list)
    preview_items = page.object_list[:5] if page.number == 1 else list(notifications[:5])

    return render(
        request,
        "core/notifications.html",
        {
            "notifications": page.object_list,
            "page_obj": page,
            "notification_summary": {"unread_count": unread_count, "items": preview_items},
            "selected_notification": (
                selected_notification
            ),
            "unread_count": (
                unread_count
            ),
        },
    )


@login_required
@require_POST
def open_notification(
    request,
    notification_id,
):
    notification = get_object_or_404(
        Notification,
        id=notification_id,
        user=request.user,
    )

    # Mark only the notification that was opened.
    notification.mark_read()

    # The target route performs its own authorization check. Legacy unsafe or
    # missing targets stay in the notification centre for their full detail.
    if notification.safe_target_url:
        return redirect(notification.safe_target_url)

    notifications_url = reverse(
        "notifications"
    )

    return redirect(
        f"{notifications_url}"
        f"?selected={notification.id}"
    )


@login_required
@require_POST
def mark_all_notifications_read(
    request,
):
    Notification.objects.filter(
        user=request.user,
        is_read=False,
    ).update(
        is_read=True,
        read_at=timezone.now(),
    )

    selected_id = request.POST.get(
        "selected"
    )

    notifications_url = reverse(
        "notifications"
    )

    params = []
    page_number = request.POST.get("page", "")
    if page_number.isdigit() and int(page_number) > 1:
        params.append(f"page={int(page_number)}")
    if selected_id and str(selected_id).isdigit():
        params.append(f"selected={int(selected_id)}")
    return redirect(f"{notifications_url}?{'&'.join(params)}" if params else notifications_url)

@login_required
def submit_story_view(request):
    profile = get_object_or_404(
        Profile,
        user=request.user
    )

    if request.method == "POST":
        form = StoryForm(request.POST)

        if form.is_valid():
            story = form.save(commit=False)
            story.user = request.user
            story.display_name = request.user.get_full_name()
            projection = get_profile_projection(profile)
            story.qualification = projection['qualification']
            story.institution = projection['institution']
            story.status = "pending"

            story.save()

            messages.success(
                request,
                "Thank you for sharing your experience! "
                "Your story has been submitted and will appear once it has been reviewed."
            )

            return redirect("stories")
        
    else:
        form = StoryForm()

    return render(request, "core/submit_story.html", {"form": form})




@login_required
def resources_view(request):

    profile = get_object_or_404(Profile, user=request.user)
    return render(request, "core/resources.html", {"profile": profile})

@login_required
def settings_view(request):
    from core.forms import NotificationPreferenceForm
    from core.models import NotificationPreference
    preference = NotificationPreference.objects.filter(user=request.user).first()
    form = NotificationPreferenceForm(request.POST if request.method == 'POST' else None,
        instance=preference or NotificationPreference(user=request.user))
    if request.method == 'POST' and form.is_valid():
        NotificationPreference.objects.update_or_create(user=request.user, defaults=form.cleaned_data)
        messages.success(request, 'Your in-app notification preferences have been saved.')
        return redirect('settings')
    return render(request, "core/settings.html", {'notification_form': form})


def _build_docx(cv_data, filepath, *, json_resume=None, layout_settings=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.shared import Pt, Inches

    resume = json_resume if isinstance(json_resume, dict) else {}
    layout = normalize_layout_settings(layout_settings)
    hidden = set(layout["hidden_sections"])
    doc = Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10.5)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.line_spacing = 1.0

    # ── Header: name centered, bold, larger; contact line centered below ──
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(cv_data.get("full_name", "CV").upper())
    name_run.bold = True
    name_run.font.size = Pt(18)
    name_p.paragraph_format.space_after = Pt(2)

    contact_bits = [b for b in [cv_data.get("email"), cv_data.get("phone"), cv_data.get("location")] if b]
    if contact_bits:
        contact_p = doc.add_paragraph(" | ".join(contact_bits))
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_p.runs[0].font.size = Pt(9)
        contact_p.paragraph_format.space_after = Pt(14)

    def add_section_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(11)

    def add_entry_header(title_org, dates):
        # Title/org on the left, dates right-aligned on the SAME line via a right tab stop —
        # this achieves the same-line layout without a table or a visible divider/border.
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT)
        title_run = p.add_run(title_org)
        title_run.bold = True
        title_run.font.size = Pt(10.5)
        if dates:
            p.add_run("\t")
            date_run = p.add_run(dates)
            date_run.italic = True
            date_run.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)

    if cv_data.get("summary") and "summary" not in hidden:
        add_section_heading("Profile")
        summary_p = doc.add_paragraph(cv_data["summary"])
        summary_p.paragraph_format.space_after = Pt(6)
        summary_p.paragraph_format.line_spacing = 1.0

    if cv_data.get("key_projects") and "projects" not in hidden:
        add_section_heading("Key Projects")
        for proj in cv_data["key_projects"]:
            add_entry_header(f"{proj.get('title', '')} | {proj.get('organisation', '')}", proj.get("dates", ""))
            if proj.get("description"):
                doc.add_paragraph(proj["description"]).paragraph_format.space_after = Pt(2)
            for bullet in proj.get("bullets", []):
                bp = doc.add_paragraph(bullet, style="List Bullet")
                bp.paragraph_format.space_after = Pt(2)

    if cv_data.get("work_experience") and "work" not in hidden:
        add_section_heading("Work Experience")
        for job in cv_data["work_experience"]:
            add_entry_header(f"{job.get('title', '')} | {job.get('organisation', '')}", job.get("dates", ""))
            for bullet in job.get("bullets", []):
                bp = doc.add_paragraph(bullet, style="List Bullet")
                bp.paragraph_format.space_after = Pt(2)
                bp.paragraph_format.line_spacing = 1.0

    if cv_data.get("education") and "education" not in hidden:
        add_section_heading("Education")
        for edu in cv_data["education"]:
            add_entry_header(f"{edu.get('qualification', '')} | {edu.get('institution', '')}", edu.get("dates", ""))

    if cv_data.get("relevant_coursework"):
        add_section_heading("Relevant Coursework")
        for group in cv_data["relevant_coursework"]:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(f"{group.get('category', '')}: ")
            r.bold = True
            p.add_run(" | ".join(group.get("items", [])))
            p.paragraph_format.space_after = Pt(2)

    if (resume.get("certificates") or cv_data.get("certifications")) and "certificates" not in hidden:
        add_section_heading("Certifications")
        certificates = resume.get("certificates") or cv_data.get("certifications") or []
        for index, cert in enumerate(certificates):
            legacy = (cv_data.get("certifications") or [])
            legacy = legacy[index] if index < len(legacy) else {}
            title = cert.get("display") or " - ".join(
                bit for bit in (cert.get("name", ""), cert.get("issuer", "")) if bit
            ) or legacy.get("title", "")
            add_entry_header(title, legacy.get("dates", "") or cert.get("date", ""))

    if cv_data.get("technical_skills") and "skills" not in hidden:
        add_section_heading("Technical Skills")
        for group in cv_data["technical_skills"]:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(f"{group.get('category', '')}: ")
            r.bold = True
            p.add_run(" | ".join(group.get("items", [])))
            p.paragraph_format.space_after = Pt(2)

    if cv_data.get("soft_skills") and "skills" not in hidden:
        add_section_heading("Soft Skills")
        doc.add_paragraph(" | ".join(cv_data["soft_skills"]))

    if resume.get("languages") and "languages" not in hidden:
        add_section_heading("Languages")
        for item in resume["languages"]:
            value = item.get("language", "")
            if item.get("fluency"):
                value = f"{value} - {item['fluency']}"
            doc.add_paragraph(value, style="List Bullet")

    if resume.get("awards") and "awards" not in hidden:
        add_section_heading("Awards")
        for item in resume["awards"]:
            add_entry_header(
                " | ".join(bit for bit in (item.get("title", ""), item.get("awarder", "")) if bit),
                item.get("date", ""),
            )
            if item.get("summary"):
                doc.add_paragraph(item["summary"])

    if resume.get("volunteer") and "volunteer" not in hidden:
        add_section_heading("Volunteering")
        for item in resume["volunteer"]:
            add_entry_header(
                " | ".join(bit for bit in (item.get("position", ""), item.get("organization", "")) if bit),
                "",
            )
            if item.get("summary"):
                doc.add_paragraph(item["summary"])

    doc.save(filepath)

def _build_pdf(cv_data, filepath):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    styles = getSampleStyleSheet()

    name_style = ParagraphStyle("CVName", parent=styles["Title"], fontSize=17, alignment=1, spaceAfter=2, leading=20)
    contact_style = ParagraphStyle("CVContact", parent=styles["Normal"], fontSize=9, alignment=1, spaceAfter=14, leading=11)
    heading_style = ParagraphStyle("CVHeading", parent=styles["Heading2"], fontSize=11, spaceBefore=10, spaceAfter=4, leading=13)
    entry_title_style = ParagraphStyle("CVEntryTitle", parent=styles["Normal"], fontSize=10.5, fontName="Helvetica-Bold", spaceAfter=1, leading=13)
    entry_dates_style = ParagraphStyle("CVEntryDates", parent=styles["Normal"], fontSize=10, fontName="Helvetica-Oblique", alignment=2, spaceAfter=3, leading=12)
    body_style = ParagraphStyle("CVBody", parent=styles["BodyText"], fontSize=10, leading=13, spaceAfter=3)
    bullet_style = ParagraphStyle("CVBullet", parent=body_style, spaceAfter=2)

    doc = SimpleDocTemplate(filepath, pagesize=A4, topMargin=1.8 * cm, bottomMargin=1.8 * cm)
    story = []

    # ── Header: name centered, bold, larger; contact line centered below ──
    story.append(Paragraph(cv_data.get("full_name", "CV").upper(), name_style))
    contact_bits = [b for b in [cv_data.get("email"), cv_data.get("phone"), cv_data.get("location")] if b]
    if contact_bits:
        story.append(Paragraph(" | ".join(contact_bits), contact_style))

    def add_entry(title_org, dates):
        story.append(Paragraph(title_org, entry_title_style))
        if dates:
            story.append(Paragraph(dates, entry_dates_style))

    if cv_data.get("summary"):
        story.append(Paragraph("Profile", heading_style))
        story.append(Paragraph(cv_data["summary"], body_style))

    if cv_data.get("key_projects"):
        story.append(Paragraph("Key Projects", heading_style))
        for proj in cv_data["key_projects"]:
            add_entry(f"{proj.get('title', '')} | {proj.get('organisation', '')}", proj.get("dates", ""))
            if proj.get("description"):
                story.append(Paragraph(proj["description"], body_style))
            for bullet in proj.get("bullets", []):
                story.append(Paragraph(f"• {bullet}", bullet_style))
            story.append(Spacer(1, 4))

    if cv_data.get("work_experience"):
        story.append(Paragraph("Work Experience", heading_style))
        for job in cv_data["work_experience"]:
            add_entry(f"{job.get('title', '')} | {job.get('organisation', '')}", job.get("dates", ""))
            for bullet in job.get("bullets", []):
                story.append(Paragraph(f"• {bullet}", bullet_style))
            story.append(Spacer(1, 4))

    if cv_data.get("education"):
        story.append(Paragraph("Education", heading_style))
        for edu in cv_data["education"]:
            add_entry(f"{edu.get('qualification', '')} | {edu.get('institution', '')}", edu.get("dates", ""))

    if cv_data.get("relevant_coursework"):
        story.append(Paragraph("Relevant Coursework", heading_style))
        for group in cv_data["relevant_coursework"]:
            story.append(Paragraph(f"<b>{group.get('category', '')}:</b> {' | '.join(group.get('items', []))}", bullet_style))

    if cv_data.get("certifications"):
        story.append(Paragraph("Certifications", heading_style))
        for cert in cv_data["certifications"]:
            add_entry(cert.get("title", ""), cert.get("dates", ""))

    if cv_data.get("technical_skills"):
        story.append(Paragraph("Technical Skills", heading_style))
        for group in cv_data["technical_skills"]:
            story.append(Paragraph(f"<b>{group.get('category', '')}:</b> {' | '.join(group.get('items', []))}", bullet_style))

    if cv_data.get("soft_skills"):
        story.append(Paragraph("Soft Skills", heading_style))
        story.append(Paragraph(" | ".join(cv_data["soft_skills"]), body_style))

    doc.build(story)




def _build_cover_letter_docx(letter_data, filepath):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(8)
    normal_style.paragraph_format.line_spacing = 1.15

    if letter_data.get("date"):
        doc.add_paragraph(letter_data["date"])
    doc.add_paragraph(letter_data.get("recipient_line", "Hiring Manager"))
    if letter_data.get("location"):
        doc.add_paragraph(letter_data["location"])
    doc.add_paragraph(letter_data.get("salutation", "Dear Hiring Manager,"))
    doc.add_paragraph(letter_data.get("opening", ""))

    for para in letter_data.get("body_paragraphs", []):
        doc.add_paragraph(para)

    doc.add_paragraph(letter_data.get("closing", ""))
    doc.add_paragraph("")
    doc.add_paragraph(letter_data.get("sign_off", "Kind regards,"))
    doc.add_paragraph(letter_data.get("signature_name", ""))

    doc.save(filepath)


def _build_cover_letter_pdf(letter_data, filepath):
    from html import escape
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    styles = getSampleStyleSheet()
    body_style = ParagraphStyle("CLBody", parent=styles["Normal"], fontSize=11, leading=15, spaceAfter=10)

    doc = SimpleDocTemplate(filepath, pagesize=A4, topMargin=2.2 * cm, bottomMargin=2.2 * cm)
    story = []

    def safe(value):
        return escape(str(value or ""))

    if letter_data.get("date"):
        story.append(Paragraph(safe(letter_data["date"]), body_style))
    story.append(Paragraph(safe(letter_data.get("recipient_line", "Hiring Manager")), body_style))
    if letter_data.get("location"):
        story.append(Paragraph(safe(letter_data["location"]), body_style))
    story.append(Paragraph(safe(letter_data.get("salutation", "Dear Hiring Manager,")), body_style))
    story.append(Paragraph(safe(letter_data.get("opening", "")), body_style))

    for para in letter_data.get("body_paragraphs", []):
        story.append(Paragraph(safe(para), body_style))

    story.append(Paragraph(safe(letter_data.get("closing", "")), body_style))
    story.append(Spacer(1, 10))
    story.append(Paragraph(safe(letter_data.get("sign_off", "Kind regards,")), body_style))
    story.append(Paragraph(safe(letter_data.get("signature_name", "")), body_style))

    doc.build(story)



@login_required
def cv_enhancer_view(request):
    profile = get_object_or_404(Profile, user=request.user)
    projection = get_profile_projection(profile)

    saved_opportunities = (
        SavedOpportunity.objects
        .filter(user=request.user, opportunity__in=Opportunity.objects.professionally_ready())
        .select_related("opportunity", "opportunity__active_clean_revision")
        .order_by("-saved_at")
    )
    selected_id = request.GET.get("opportunity", "").strip()
    preselected_opportunity = None
    if selected_id.isdigit():
        preselected_opportunity = Opportunity.objects.professionally_ready().filter(pk=int(selected_id)).first()

    skill_names = projection['skill_names']
    context = {
        "profile": profile,
        "projection": projection,
        "profile_name": request.user.get_full_name() or request.user.username,
        "skill_names": skill_names,
        "skills_count": len(skill_names),
        "saved_opportunities": saved_opportunities,
        "saved_count": saved_opportunities.count(),
        "cv_count": GeneratedDocument.objects.filter(user=request.user, document_type="cv").count(),
        "selected_opportunity_id": int(selected_id) if selected_id.isdigit() else None,
        "preselected_opportunity": preselected_opportunity,
        "preselected_is_saved": bool(preselected_opportunity and saved_opportunities.filter(opportunity=preselected_opportunity).exists()),
        "resume_templates": public_template_options(),
        "default_resume_template": "azurill",
    }
    return render(request, "core/cv_enhancer.html", context)


@login_required
def cover_letter_generator_view(request):
    profile = get_object_or_404(Profile, user=request.user)
    projection = get_profile_projection(profile)

    saved_opportunities = (
        SavedOpportunity.objects
        .filter(user=request.user, opportunity__in=Opportunity.objects.professionally_ready())
        .select_related("opportunity", "opportunity__active_clean_revision")
        .order_by("-saved_at")
    )
    selected_id = request.GET.get("opportunity", "").strip()
    preselected_opportunity = None
    if selected_id.isdigit():
        preselected_opportunity = Opportunity.objects.professionally_ready().filter(pk=int(selected_id)).first()

    skill_names = projection['skill_names']
    context = {
        "profile": profile,
        "projection": projection,
        "profile_name": request.user.get_full_name() or request.user.username,
        "skill_names": skill_names,
        "skills_count": len(skill_names),
        "saved_opportunities": saved_opportunities,
        "saved_count": saved_opportunities.count(),
        "cv_count": GeneratedDocument.objects.filter(user=request.user, document_type="cv").count(),
        "selected_opportunity_id": int(selected_id) if selected_id.isdigit() else None,
        "preselected_opportunity": preselected_opportunity,
        "preselected_is_saved": bool(preselected_opportunity and saved_opportunities.filter(opportunity=preselected_opportunity).exists()),
    }
    return render(request, "core/cover_letter_generator.html", context)


# ------------------------------------------------------------
# UPDATED: CV prompt builder
# ------------------------------------------------------------

def _build_cv_prompt(profile, job_description, cv_text="", additional_notes=""):
    cv_block = ""
    if cv_text:
        cv_block = f"""
The candidate has provided their existing CV text below. Use it as the
primary source of truth for their experience, projects, and education —
enhance and restructure it rather than starting from scratch:
\"\"\"{cv_text}\"\"\"
"""

    notes_block = ""
    if additional_notes:
        notes_block = f"""
The candidate specifically asked you to highlight or take into account
the following when writing the CV:
\"\"\"{additional_notes}\"\"\"
"""

    return f"""
You are an elite CV writer specialising in repositioning candidates with little or
unclear experience into stronger, employable professionals. Your objective is to
transform raw or average CV content into results driven, comercially valuable cv
aligned with SA hiring standards following these strict rules, the summary must be objective
(no descriptive words like 'motivated', 'hardworking', etc) every bullet point must show atleast
one of the following: Outcomes, Deliverable, responsibility, measurable impact where possible.
Remove all fluff and generic statements. Prioritise clarity, structure, and value over length (1 page cv).
Write like a candidate has already operated in a professional environment. Using the candidate's profile below and the
job requirements provided, write a tailored CV. Respond ONLY with valid JSON,
no markdown fences, no commentary, matching exactly this shape:

{{
  "full_name": "string",
  "summary": "3-4 sentence professional summary tailored to the role",
  "key_projects": [
    {{
      "title": "string",
      "organisation": "string",
      "dates": "string",
      "description": "one sentence describing the project",
      "bullets": ["achievement 1", "achievement 2"]
    }}
  ],
  "work_experience": [
    {{
      "title": "string",
      "organisation": "string",
      "dates": "string",
      "bullets": ["achievement 1", "achievement 2"]
    }}
  ],
  "education": [
    {{"qualification": "string", "institution": "string", "dates": "string"}}
  ],
  "relevant_coursework": [
    {{"category": "string", "items": ["item1", "item2"]}}
  ],
  "certifications": [
    {{"title": "string", "dates": "string"}}
  ],
  "technical_skills": [
    {{"category": "string", "items": ["item1", "item2"]}}
  ],
  "soft_skills": ["skill1", "skill2"]
}}

Only include keys that have real content — omit "key_projects" entirely if the
candidate has no relevant projects, rather than inventing filler entries.

Candidate profile:
Name: {profile.get('full_name', '')}
Qualification: {profile.get('qualification', '')}
Institution: {profile.get('institution', '')}
Skills: {', '.join(profile.get('skills', []))}
Location: {profile.get('location', '')}
{cv_block}{notes_block}
Job requirements pasted by the candidate:
\"\"\"{job_description}\"\"\"
"""


def _generate_cv_content(profile, job_description, cv_text="", additional_notes=""):
    from google import genai

    client = genai.Client(api_key=settings.GEMINI_API_KEY)
    response = client.models.generate_content(
        model="gemini-3.5-flash",
        contents=_build_cv_prompt(profile, job_description, cv_text, additional_notes),
    )
    text = response.text.strip()
    if text.startswith("```"):
        text = text.strip("`")
        text = text.split("\n", 1)[1] if "\n" in text else text
        text = text.rsplit("```", 1)[0]
    return json.loads(text)


@login_required
@require_POST
@document_timing("cv")
def generate_cv(request):
    manual_job_description = (
        request.POST.get("job_description", "").strip()
        or request.POST.get("job_requirements", "").strip()
    )
    opportunity_id = request.POST.get("opportunity_id", "").strip()
    opportunity = None
    application = None

    if opportunity_id:
        if not opportunity_id.isdigit():
            return JsonResponse({"error": "Invalid opportunity."}, status=400)
        opportunity = Opportunity.objects.professionally_ready().select_related(
            "active_clean_revision").filter(pk=int(opportunity_id)).first()
        if opportunity is None:
            return JsonResponse({"error": "That opportunity is no longer available in CareerGPS."}, status=404)
        allowed = (
            SavedOpportunity.objects.filter(user=request.user, opportunity=opportunity).exists()
            or Application.objects.filter(user=request.user, opportunity=opportunity).exists()
        )
        if not allowed:
            return JsonResponse({
                "error": "Select a saved opportunity or open it from your application workspace."
            }, status=403)
        with timed_stage("persistence"):
            application, _ = Application.objects.get_or_create(
                user=request.user, opportunity=opportunity,
                defaults={"clean_revision": opportunity.active_clean_revision})
            if not application.clean_revision_id:
                application.clean_revision = opportunity.active_clean_revision
        with timed_stage("job_context"):
            job_context = resolve_job_context(
                opportunity, clean_revision=application.clean_revision)
        job_data = job_context["data"]
        with timed_stage("persistence"):
            application.job_info_status = job_context["status"]
            application.save(update_fields=["clean_revision", "job_info_status", "updated_at"])
    else:
        if not manual_job_description:
            return JsonResponse({"error": "Please paste the job requirements first, or select a saved opportunity."}, status=400)
        job_context = {"status": "manual", "limited": False}
        job_data = {
            "title": "", "employer": "", "location": "",
            "description": manual_job_description,
            "qualifications": [], "experience": [], "skills": [],
            "certifications": [], "responsibilities": [],
            "preferred_requirements": [], "other_requirements": [],
            "deadline": "", "work_arrangement": "", "employment_type": "",
            "application_instructions": "", "source_url": "",
        }

    profile_obj = getattr(request.user, "profile", None)
    if profile_obj is None:
        return JsonResponse({"error": "Please complete your profile before generating a CV."}, status=400)

    cv_text = request.POST.get("cv_text", "").strip()
    cv_file = request.FILES.get("cv_file")
    if not cv_text and cv_file:
        max_size = 5 * 1024 * 1024
        if cv_file.size > max_size:
            return JsonResponse({"error": "That file is too large. Please upload a file under 5MB."}, status=400)
        try:
            cv_text = _extract_text_from_cv(cv_file)
        except ValueError as exc:
            return JsonResponse({"error": str(exc)}, status=400)

    additional_notes = request.POST.get("additional_notes", "").strip()
    template_key = normalize_template_key(request.POST.get("template_key"))
    layout_settings = normalize_layout_settings({"density": request.POST.get("density", "comfortable")})

    try:
        with timed_stage("candidate_bundle"):
            candidate_bundle = build_candidate_fact_bundle(profile_obj, raw_cv_text=cv_text)
        with timed_stage("job_context"):
            job_text = job_context_text(job_data)
        with timed_stage("semantic_selection"):
            selected_facts = select_relevant_facts(candidate_bundle, job_text)
        selected_ids = [fact["id"] for fact in selected_facts]
        json_resume = build_json_resume(candidate_bundle, selected_fact_ids=selected_ids)
        with timed_stage("gemini"):
            writing = generate_cv_writing(
                candidate_bundle, selected_facts, job_data, additional_notes=additional_notes
            )
        cv_data = assemble_legacy_cv(candidate_bundle, json_resume, writing)
    except Exception as exc:
        return JsonResponse({"error": f"Generation failed: {exc}"}, status=500)

    with timed_stage("persistence"):
        file_id = uuid.uuid4().hex
        out_dir = os.path.join(settings.MEDIA_ROOT, "generated_cvs")
        os.makedirs(out_dir, exist_ok=True)
        docx_path = os.path.join(out_dir, f"{file_id}.docx")
        pdf_path = os.path.join(out_dir, f"{file_id}.pdf")
    with timed_stage("docx"):
        _build_docx(
            cv_data,
            docx_path,
            json_resume=json_resume,
            layout_settings=layout_settings,
        )
    with timed_stage("pdf"):
        template_key, layout_settings = render_resume_pdf(
            json_resume, cv_data, pdf_path, template_key=template_key, layout_settings=layout_settings
        )

    with timed_stage("persistence"):
        document = GeneratedDocument.objects.create(
            user=request.user,
            application=application,
            opportunity=opportunity,
            clean_revision=application.clean_revision if application else None,
            document_type="cv",
            structured_data=cv_data,
            json_resume=json_resume,
            candidate_fact_ids=selected_ids,
            job_info_status=job_context["status"],
            template_key=template_key,
            layout_settings=layout_settings,
            renderer_backend="careergps_rr_layout_adapter",
            pdf_file=f"generated_cvs/{file_id}.pdf",
            docx_file=f"generated_cvs/{file_id}.docx",
        )

    return JsonResponse({
        "docx_url": reverse("applications:document_download", args=[document.pk, "docx"]),
        "pdf_url": reverse("applications:document_download", args=[document.pk, "pdf"]),
        "job_info_status": job_context["status"],
        "application_id": application.id if application else None,
        "template_key": document.template_key,
        "preview_url": reverse("applications:document_preview", args=[document.id]),
        "editor_url": reverse("applications:cv_editor", args=[document.id]),
    })


# ------------------------------------------------------------
# UPDATED: Cover letter prompt builder
# ------------------------------------------------------------

TONE_INSTRUCTIONS = {
    "professional": "Keep the tone polished, measured, and businesslike.",
    "enthusiastic": "Let genuine enthusiasm for the role and organisation come through, without sounding over-the-top or generic.",
    "confident": "Write with a confident, assured tone that leads with the candidate's strengths.",
    "humble": "Keep the tone modest and grounded, letting the evidence speak for itself rather than overselling.",
    "creative": "Allow a little more personality and creative phrasing, while staying professional and credible.",
    "direct": "Keep sentences short and to the point — no filler, no throat-clearing.",
}


def _build_cover_letter_prompt(cv_text, job_description, additional_notes="", tone="professional"):
    tone_instruction = TONE_INSTRUCTIONS.get(tone, TONE_INSTRUCTIONS["professional"])

    notes_block = ""
    if additional_notes:
        notes_block = f"""
The candidate specifically asked you to highlight or take into account
the following:
\"\"\"{additional_notes}\"\"\"
"""

    return f"""
You are an expert professional cover letter writing assistant.

Your task is to create a highly tailored cover letter using ONLY:
1. The candidate's CV text.
2. The job/opportunity requirements provided below.

TONE: {tone_instruction}

The cover letter must feel like it was written specifically for this
candidate and this employer. It must NOT simply repeat the candidate's CV.

IMPORTANT RULES:
- Base every factual claim strictly on the CV and job information provided.
- NEVER invent employers, responsibilities, achievements, qualifications,
  skills, projects, statistics, company facts, or experiences.
- Do not assume information that is not explicitly provided.
- If the job information does not provide a company name, contact name,
  company-specific detail, or specific business problem, do not invent one.
- Use natural, confident, professional language.
- Avoid generic, robotic, overly enthusiastic, or AI-sounding language.
- Do not use phrases such as "I am writing to apply for..."
  or "I am excited to apply for this position..."
- Do not turn the cover letter into another CV.
- Do not simply list qualifications, skills, or work experience.
- Every paragraph should have a clear purpose.

COVER LETTER STRUCTURE:

1. OPENING / HOOK
Start with a strong hook that demonstrates why the candidate is relevant
to the opportunity.

The opening should lead with one of the following, where supported by the
CV:
- A relevant result or achievement.
- A meaningful experience.
- A relevant technical/business insight.
- A problem the candidate has experience solving.
- A strength that directly relates to what the employer needs.

Do NOT begin with:
"I am writing to apply for..."
"I would like to apply for..."
"I am excited to apply for..."

The opening should naturally establish the role/opportunity without making
the fact that they are applying the main point of the paragraph.

2. THE STORY BEHIND THE CV
Do not repeat the CV as a list.

Instead, explain the connection between the candidate's experiences and
the opportunity. Show how their education, projects, work experience,
internships, certifications, or other relevant experiences have prepared
them to contribute to this particular role.

Focus on the meaning and relevance of their experience rather than merely
stating that they have it.

3. EMPLOYER-FOCUSED VALUE
The cover letter must focus on what the candidate can contribute to the
employer.

Identify the employer's needs from the job requirements and connect them
directly to evidence from the candidate's CV.

Use the structure:

Employer need → Candidate evidence → Potential contribution

Do not focus heavily on statements such as:
"I want to grow my skills."
"This role will help me achieve my goals."
"I want to gain experience."

The employer should clearly understand what the candidate can bring to
the organisation.

4. WHY THIS COMPANY / OPPORTUNITY
Where the job information provides enough evidence, explain why this
specific company, organisation, role, or opportunity is relevant to the
candidate.

Do not use generic statements that could apply to any company.

If the job description contains something specific about:
- the company's work,
- its products or services,
- the role,
- its technology,
- its business challenges,
- its values,
- its industry,
- or the responsibilities,

connect that detail to something genuinely supported by the candidate's CV.

If no company-specific information is provided, do NOT invent company
details. Instead, focus on the specific requirements and responsibilities
of the opportunity.

5. PERSONALISATION
Personalise the letter using at least three relevant elements whenever
the provided information allows it:

A. The company/organisation name.
B. One specific detail about the company's work, role, responsibilities,
   requirements, or opportunity.
C. One direct connection between the candidate's actual experience and
   that specific employer need.

The personalisation must feel natural rather than forced.

6. UNIQUE VALUE
Identify what makes this candidate particularly relevant to the role based
on the evidence available in the CV.

This could be a combination of:
- technical skills,
- business knowledge,
- academic background,
- projects,
- problem-solving experience,
- communication skills,
- leadership,
- customer experience,
- certifications,
- or other demonstrated strengths.

Do not claim that the candidate is "unique" unless the evidence supports
the statement. Instead, demonstrate their value through the connection
between their experience and the employer's needs.

7. CLOSING
End confidently and professionally.

Do NOT use:
"I hope to hear from you soon."
"Thank you for considering my application and I hope to hear from you."

Instead, use a closing that communicates the candidate's interest in
contributing to the organisation and invites a conversation or next step.

For example, the closing should communicate the idea:
"I'd welcome the opportunity to discuss how my experience could contribute
to [Company]."

Do not copy this example word-for-word every time.

OUTPUT REQUIREMENTS:

Respond ONLY with valid JSON.
Do NOT include markdown fences.
Do NOT include commentary outside the JSON.
Return exactly this structure:

{{
  "recipient_line": " Dear Recruiter Team",
  "opening": "A strong 1-2 sentence hook that leads with relevance, insight, achievement, experience, or value rather than the intention to apply.",
  "body_paragraphs": [
    "Paragraph explaining the story behind the candidate's CV and how their relevant experience connects to the opportunity.",
    "Paragraph connecting the employer's specific needs to evidence from the candidate's CV and explaining the value the candidate can contribute.",
    "Optional paragraph explaining why this specific company, organisation, role, or opportunity is relevant, using specific information from the job requirements."
  ],
  "closing": "A confident 1-2 sentence closing that reinforces the candidate's potential contribution and invites a conversation or next step.",
  "signature_name": "The candidate's full name exactly as found in the CV."
}}

ADDITIONAL QUALITY CHECK BEFORE RESPONDING:

Before generating the final JSON, silently check that:

1. The opening does NOT begin with "I am writing to apply".
2. The opening contains a genuine hook based on available evidence.
3. The letter does not simply repeat the CV.
4. The letter explains the story or connection behind the candidate's
   experience.
5. The employer's needs are prioritised over the candidate's personal goals.
6. The letter clearly explains what the candidate can contribute.
7. The company/opportunity is personalised where information is available.
8. At least three personalisation elements are used when the provided
   information makes this possible.
9. Every claim can be traced to the CV or job requirements.
10. No information has been invented.
11. The closing invites a conversation or next step rather than saying
    "I hope to hear from you soon."
12. The language sounds like a real professional candidate, not a generic
    AI-generated template.

Candidate's CV text:
\"\"\"{cv_text}\"\"\"
{notes_block}
Job requirements pasted by the candidate:
\"\"\"{job_description}\"\"\"

"""


def _generate_cover_letter_content(cv_text, job_description, additional_notes="", tone="professional"):
    from google import genai

    client = genai.Client(api_key=settings.GEMINI_API_KEY)
    response = client.models.generate_content(
        model="gemini-3.5-flash",
        contents=_build_cover_letter_prompt(cv_text, job_description, additional_notes, tone),
    )
    text = response.text.strip()
    if text.startswith("```"):
        text = text.strip("`")
        text = text.split("\n", 1)[1] if "\n" in text else text
        text = text.rsplit("```", 1)[0]
    return json.loads(text)


@login_required
@require_POST
@document_timing("cover_letter")
def generate_cover_letter(request):
    manual_job_description = (
        request.POST.get("job_description", "").strip()
        or request.POST.get("job_requirements", "").strip()
    )
    opportunity_id = request.POST.get("opportunity_id", "").strip()
    opportunity = None
    application = None

    if opportunity_id:
        if not opportunity_id.isdigit():
            return JsonResponse({"error": "Invalid opportunity."}, status=400)
        opportunity = Opportunity.objects.professionally_ready().select_related(
            "active_clean_revision").filter(pk=int(opportunity_id)).first()
        if opportunity is None:
            return JsonResponse({"error": "That opportunity is no longer available in CareerGPS."}, status=404)
        allowed = (
            SavedOpportunity.objects.filter(user=request.user, opportunity=opportunity).exists()
            or Application.objects.filter(user=request.user, opportunity=opportunity).exists()
        )
        if not allowed:
            return JsonResponse({
                "error": "Select a saved opportunity or open it from your application workspace."
            }, status=403)
        with timed_stage("persistence"):
            application, _ = Application.objects.get_or_create(
                user=request.user, opportunity=opportunity,
                defaults={"clean_revision": opportunity.active_clean_revision})
            if not application.clean_revision_id:
                application.clean_revision = opportunity.active_clean_revision
        with timed_stage("job_context"):
            job_context = resolve_job_context(
                opportunity, clean_revision=application.clean_revision)
        job_data = job_context["data"]
        with timed_stage("persistence"):
            application.job_info_status = job_context["status"]
            application.save(update_fields=["clean_revision", "job_info_status", "updated_at"])
    else:
        if not manual_job_description:
            return JsonResponse({"error": "Please paste the job requirements first, or select a saved opportunity."}, status=400)
        job_context = {"status": "manual", "limited": False}
        job_data = {
            "title": "", "employer": "", "location": "",
            "description": manual_job_description,
            "qualifications": [], "experience": [], "skills": [],
            "certifications": [], "responsibilities": [],
            "preferred_requirements": [], "other_requirements": [],
            "deadline": "", "work_arrangement": "", "employment_type": "",
            "application_instructions": "", "source_url": "",
        }

    profile_obj = getattr(request.user, "profile", None)
    if profile_obj is None:
        return JsonResponse({"error": "Please complete your profile before generating a cover letter."}, status=400)

    cv_text = request.POST.get("cv_text", "").strip()
    cv_file = request.FILES.get("cv_file")
    if not cv_text and cv_file:
        max_size = 5 * 1024 * 1024
        if cv_file.size > max_size:
            return JsonResponse({"error": "That file is too large. Please upload a file under 5MB."}, status=400)
        try:
            cv_text = _extract_text_from_cv(cv_file)
        except ValueError as exc:
            return JsonResponse({"error": str(exc)}, status=400)

    additional_notes = request.POST.get("additional_notes", "").strip()
    tone = request.POST.get("tone", "professional").strip() or "professional"

    try:
        with timed_stage("candidate_bundle"):
            candidate_bundle = build_candidate_fact_bundle(profile_obj, raw_cv_text=cv_text)
        with timed_stage("job_context"):
            job_text = job_context_text(job_data)
        with timed_stage("semantic_selection"):
            selected_facts = select_relevant_facts(candidate_bundle, job_text)
        selected_ids = [fact["id"] for fact in selected_facts]
        with timed_stage("gemini"):
            result = generate_controlled_cover_letter(
                candidate_bundle, selected_facts, job_data,
                additional_notes=additional_notes, tone=tone,
            )
        letter_data = result.model_dump(mode="json")
        recipient_bits = [letter_data.get("recipient_name", "Hiring Manager"), letter_data.get("company_name", "")]
        letter_data["recipient_line"] = ": ".join(bit for bit in recipient_bits if bit)
    except Exception as exc:
        return JsonResponse({"error": f"Generation failed: {exc}"}, status=500)

    with timed_stage("persistence"):
        file_id = uuid.uuid4().hex
        out_dir = os.path.join(settings.MEDIA_ROOT, "generated_cover_letters")
        os.makedirs(out_dir, exist_ok=True)
        docx_path = os.path.join(out_dir, f"{file_id}.docx")
        pdf_path = os.path.join(out_dir, f"{file_id}.pdf")
    with timed_stage("docx"):
        _build_cover_letter_docx(letter_data, docx_path)
    with timed_stage("pdf"):
        _build_cover_letter_pdf(letter_data, pdf_path)

    with timed_stage("persistence"):
        document = GeneratedDocument.objects.create(
            user=request.user,
            application=application,
            opportunity=opportunity,
            clean_revision=application.clean_revision if application else None,
            document_type="cover_letter",
            structured_data=letter_data,
            candidate_fact_ids=selected_ids,
            job_info_status=job_context["status"],
            template_key="careergps_professional",
            renderer_backend="django_legacy",
            pdf_file=f"generated_cover_letters/{file_id}.pdf",
            docx_file=f"generated_cover_letters/{file_id}.docx",
        )

    return JsonResponse({
        "docx_url": reverse("applications:document_download", args=[document.pk, "docx"]),
        "pdf_url": reverse("applications:document_download", args=[document.pk, "pdf"]),
        "job_info_status": job_context["status"],
        "application_id": application.id if application else None,
        "preview_url": reverse("applications:document_preview", args=[document.id]),
        "editor_url": reverse("applications:document_edit", args=[document.id]),
    })
