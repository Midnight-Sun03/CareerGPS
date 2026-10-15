"""Private vacancy intake: source text is untrusted until employer confirmation."""
from datetime import datetime
from io import BytesIO
import re
from zipfile import ZipFile

from django.conf import settings
from django.core.exceptions import ValidationError

from ingestion.ai.service import _criterion_rows, _evidence_supported

MAX_UPLOAD = 5 * 1024 * 1024
MAX_TEXT = 50000


def read_vacancy_document(upload):
    name = str(upload.name or '').rsplit('/', 1)[-1].rsplit('\\', 1)[-1]
    extension = name.lower().rsplit('.', 1)[-1] if '.' in name else ''
    if extension not in {'pdf', 'docx', 'txt'}:
        raise ValidationError('Upload a PDF, DOCX or plain-text vacancy, or paste its text.')
    data = upload.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise ValidationError('The vacancy document must be no larger than 5 MiB.')
    try:
        if extension == 'txt':
            if b'\x00' in data:
                raise ValueError('binary text')
            text = data.decode('utf-8-sig')
        elif extension == 'pdf':
            if not data.startswith(b'%PDF-'):
                raise ValueError('invalid PDF')
            import pdfplumber
            with pdfplumber.open(BytesIO(data)) as pdf:
                if len(pdf.pages) > 50:
                    raise ValueError('too many pages')
                text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        else:
            if not data.startswith(b'PK'):
                raise ValueError('invalid DOCX')
            with ZipFile(BytesIO(data)) as archive:
                files = archive.infolist()
                if len(files) > 1000 or sum(item.file_size for item in files) > 20 * 1024 * 1024:
                    raise ValueError('oversized document contents')
            from docx import Document
            document = Document(BytesIO(data))
            text = '\n'.join([*(p.text for p in document.paragraphs),
                *(cell.text for table in document.tables for row in table.rows for cell in row.cells)])
    except Exception as exc:
        raise ValidationError('Could not safely read text from this document. Paste the vacancy text or use manual entry.') from exc
    return validate_source_text(text), name


def validate_source_text(value):
    text = str(value or '').strip()
    if len(text) < 30:
        raise ValidationError('No usable vacancy text was found. Paste the vacancy text or use manual entry.')
    if len(text) > MAX_TEXT:
        raise ValidationError('Vacancy text is too long. Provide at most 50,000 characters.')
    return text


def _facts(text):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    facts = {'description': text}
    labelled = {
        'title': r'^(?:job title|position|role)\s*:\s*(.+)$',
        'location': r'^(?:location|work location)\s*:\s*(.+)$',
        'employment_type': r'^(?:employment type|contract type)\s*:\s*(.+)$',
        'application_instructions': r'^(?:how to apply|application instructions)\s*:\s*(.+)$',
    }
    for line in lines:
        for field, pattern in labelled.items():
            match = re.match(pattern, line, re.I)
            if match and field not in facts:
                facts[field] = match.group(1).strip()[:200 if field != 'application_instructions' else 2000]
        arrangement = re.match(r'^(?:work arrangement|work mode)\s*:\s*(on.?site|hybrid|remote)\b', line, re.I)
        if arrangement:
            value = arrangement.group(1).lower().replace('-', '')
            facts['work_arrangement'] = 'onsite' if value == 'onsite' else value
        deadline = re.match(r'^(?:closing date|application deadline)\s*:\s*(\d{4}-\d{2}-\d{2})\b', line, re.I)
        if deadline:
            try:
                facts['closing_date'] = datetime.strptime(deadline.group(1), '%Y-%m-%d').date().isoformat()
            except ValueError:
                pass
    if ('title' not in facts and lines and len(lines[0]) <= 100 and ':' not in lines[0]
            and not lines[0].endswith('.') and lines[0].casefold() not in {
                'requirements', 'responsibilities', 'about us', 'job description', 'vacancy'}):
        facts['title'] = lines[0]
    return facts


def _fallback_requirements(text):
    rows = []
    section = ''
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        heading = line.rstrip(':').casefold()
        if heading in {'requirements', 'required', 'minimum requirements', 'essential requirements'}:
            section = 'required'; continue
        if heading in {'preferred', 'preferred requirements', 'nice to have', 'advantageous'}:
            section = 'preferred'; continue
        if heading in {'responsibilities', 'duties', 'about us', 'benefits', 'how to apply'}:
            section = ''; continue
        if not section or not re.match(r'^[-•*]\s+|^\d+[.)]\s+', line):
            continue
        value = re.sub(r'^([-•*]|\d+[.)])\s+', '', line).strip()
        lower = value.casefold()
        kind = ('qualification' if re.search(r'\b(degree|diploma|qualification|nqf)\b', lower) else
            'certification' if re.search(r'\b(certification|certificate|licen[cs]e|registration)\b', lower) else
            'experience' if re.search(r'\bexperience\b|\d+\s+years?', lower) else
            'language' if re.search(r'\b(language|fluent|proficient in english)\b', lower) else
            'work_eligibility' if re.search(r'\b(right to work|work permit|citizenship)\b', lower) else 'skill')
        years = re.search(r'\b(\d{1,2})\+?\s+years?\b', lower) if kind == 'experience' else None
        rows.append({'criterion_type': kind, 'requirement_level': section, 'raw_value': value,
            'minimum_years': int(years.group(1)) if years else None, 'maximum_years': None,
            'minimum_nqf': None, 'context': '', 'source_excerpt': line[:500]})
    return rows[:30]


def interpret_vacancy(text, *, extractor=None):
    """Only source-supported proposals; no result here is authoritative."""
    text = validate_source_text(text)
    facts = _facts(text)
    fallback = _fallback_requirements(text)
    method, notice, proposals = 'deterministic', '', fallback
    if extractor is None and getattr(settings, 'GEMINI_API_KEY', ''):
        from ingestion.ai.extractor import GeminiOpportunityExtractor
        extractor = GeminiOpportunityExtractor()
    if extractor is not None:
        payload = {'opportunity_id': 0, 'source_completeness': 'full_text', 'sections': {'vacancy_text': text}}
        try:
            understood = extractor.extract_one(payload)
            if understood.opportunity_id != 0:
                raise ValueError('Incorrect extraction target')
            extracted = []
            for row in _criterion_rows(understood):
                if not _evidence_supported(row['evidence_text'], payload, row['source_section']):
                    raise ValueError('Unsupported source excerpt')
                extracted.append({'criterion_type': row['criterion_type'],
                    'requirement_level': row['requirement_level'], 'raw_value': row['raw_value'],
                    'minimum_years': row['minimum_years'], 'maximum_years': row['maximum_years'],
                    'minimum_nqf': None, 'context': row['metadata'].get('context', ''),
                    'source_excerpt': row['evidence_text'][:500]})
            if extracted:
                proposals, method = extracted[:30], 'gemini_proposal'
        except Exception:
            notice = 'Language extraction was unavailable. Only directly identifiable text was proposed; review and add missing requirements manually.'
    return facts, proposals, method, notice
